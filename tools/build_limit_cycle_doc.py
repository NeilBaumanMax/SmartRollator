from __future__ import annotations

import math
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "generated_docs"
ASSET_DIR = OUT_DIR / "assets"
DOCX_PATH = OUT_DIR / "Hopf极限环理解_智能助行车.docx"
FIG_PATH = ASSET_DIR / "hopf_limit_cycle_explanation.png"


def ensure_dirs() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)


def find_font() -> str:
    for path in [
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
    ]:
        if Path(path).exists():
            return path
    return ""


def get_font(size: int):
    path = find_font()
    return ImageFont.truetype(path, size=size) if path else ImageFont.load_default()


def draw_arrow(draw, p1, p2, fill, width=4):
    draw.line([p1, p2], fill=fill, width=width)
    x1, y1 = p1
    x2, y2 = p2
    ang = math.atan2(y2 - y1, x2 - x1)
    size = 16
    pts = [
        (x2, y2),
        (x2 - size * math.cos(ang - 0.45), y2 - size * math.sin(ang - 0.45)),
        (x2 - size * math.cos(ang + 0.45), y2 - size * math.sin(ang + 0.45)),
    ]
    draw.polygon(pts, fill=fill)


def center_text(draw, box, text, font, fill=(15, 23, 42)):
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=font, spacing=6)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    draw.multiline_text(
        (x1 + (x2 - x1 - tw) / 2, y1 + (y2 - y1 - th) / 2),
        text,
        font=font,
        fill=fill,
        align="center",
        spacing=6,
    )


def create_limit_cycle_figure() -> None:
    img = Image.new("RGB", (1800, 1180), "white")
    draw = ImageDraw.Draw(img)
    title = get_font(42)
    h = get_font(29)
    body = get_font(23)
    small = get_font(18)

    draw.text((72, 46), "怎样理解 Hopf 振荡器的极限环", font=title, fill=(15, 23, 42))
    draw.text(
        (74, 104),
        "极限环不是一个点，而是一条稳定闭合轨道：状态偏小会被推向外，偏大会被拉回来，最终围绕固定半径持续旋转。",
        font=body,
        fill=(71, 85, 105),
    )

    # Phase plane panel
    panel = (70, 170, 1040, 1040)
    draw.rounded_rectangle(panel, radius=18, outline=(37, 99, 235), width=3, fill=(248, 251, 255))
    draw.text((105, 205), "1. x-y 平面中的稳定闭合轨道", font=h, fill=(30, 64, 175))
    cx, cy, r = 550, 620, 255
    draw.line((cx - 360, cy, cx + 360, cy), fill=(148, 163, 184), width=2)
    draw.line((cx, cy + 330, cx, cy - 330), fill=(148, 163, 184), width=2)
    draw.text((cx + 372, cy - 10), "x", font=body, fill=(51, 65, 85))
    draw.text((cx + 12, cy - 363), "y", font=body, fill=(51, 65, 85))

    # limit ring
    draw.ellipse((cx - r, cy - r, cx + r, cy + r), outline=(37, 99, 235), width=7)
    draw.text((cx - 85, cy + r + 34), "r = sqrt(μ)", font=body, fill=(30, 64, 175))

    # inside spiral outward
    pts = []
    for i in range(210):
        t = i / 24
        rr = 38 + 0.88 * i
        pts.append((cx + rr * math.cos(t), cy - rr * math.sin(t)))
    draw.line(pts, fill=(16, 185, 129), width=5)
    draw_arrow(draw, pts[-14], pts[-1], (16, 185, 129), width=4)
    draw.text((cx - 320, cy - 250), "初始半径小于目标\n径向项推动它向外", font=small, fill=(4, 120, 87))

    # outside spiral inward
    pts2 = []
    for i in range(210):
        t = 0.7 + i / 25
        rr = 365 - 0.54 * i
        pts2.append((cx + rr * math.cos(t), cy - rr * math.sin(t)))
    draw.line(pts2, fill=(239, 68, 68), width=5)
    draw_arrow(draw, pts2[-14], pts2[-1], (239, 68, 68), width=4)
    draw.text((cx + 185, cy - 295), "初始半径大于目标\n径向项把它拉回", font=small, fill=(185, 28, 28))

    # rotation arrow
    draw.arc((cx - r, cy - r, cx + r, cy + r), start=292, end=26, fill=(245, 158, 11), width=10)
    draw_arrow(draw, (cx + 217, cy - 136), (cx + 246, cy - 45), (245, 158, 11), width=4)
    draw.text((cx + 125, cy + 248), "ω 决定沿环转得多快", font=body, fill=(180, 83, 9))

    # Explanation boxes
    right = (1100, 170, 1730, 1040)
    draw.rounded_rectangle(right, radius=18, outline=(15, 118, 110), width=3, fill=(246, 254, 252))
    draw.text((1130, 205), "2. 和代码、助行车的关系", font=h, fill=(17, 94, 89))

    boxes = [
        ((1160, 290, 1670, 405), "代码计算当前半径平方\nr² = x² + y²"),
        ((1160, 470, 1670, 585), "和目标半径平方比较\nμ - r²"),
        ((1160, 650, 1670, 765), "更新 Hopf 状态\nx、y 持续但有界地振荡"),
        ((1160, 830, 1670, 945), "用 x 调制导纳阻尼\nD = D_base - λx"),
    ]
    for box, text in boxes:
        draw.rounded_rectangle(box, radius=14, outline=(15, 118, 110), width=2, fill=(255, 255, 255))
        center_text(draw, box, text, body, fill=(15, 78, 72))
    for a, b in zip(boxes[:-1], boxes[1:]):
        draw_arrow(draw, ((a[0][0] + a[0][2]) / 2, a[0][3]), ((b[0][0] + b[0][2]) / 2, b[0][1]), (15, 118, 110), width=4)

    draw.text(
        (1130, 990),
        "结论：极限环提供的是稳定节律，不是直接轮速。车轮速度仍由导纳模型和限幅决定。",
        font=small,
        fill=(51, 65, 85),
    )

    img.save(FIG_PATH)


def set_east_asia(run, name="Microsoft YaHei"):
    run.font.name = "Calibri"
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7DEE8", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_doc(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(17, 24, 39)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    for name, size, color, before, after in [
        ("Title", 24, "0B2545", 0, 12),
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_p(doc, text="", style=None):
    p = doc.add_paragraph(style=style)
    r = p.add_run(text)
    set_east_asia(r)
    return p


def add_formula(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F7FB")
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    set_table_borders(table, color="CBD5E1", size="8")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, line in enumerate(lines):
        if i:
            p.add_run("\n")
        r = p.add_run(line)
        r.font.name = "Cambria Math"
        r._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(13)
    doc.add_paragraph()


def add_callout(doc, title, body, fill="F4F7FB", border="93C5FD"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    table.columns[0].width = Inches(6.3)
    set_table_borders(table, color=border, size="10")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor(30, 64, 175)
    set_east_asia(r)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    set_east_asia(r2)
    doc.add_paragraph()


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    set_table_borders(table)
    for i, w in enumerate(widths):
        for cell in table.columns[i].cells:
            cell.width = Inches(w)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(30, 64, 175)
        set_east_asia(r)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            r.font.size = Pt(9.3)
            set_east_asia(r)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_east_asia(r)


def add_numbers(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(item)
        set_east_asia(r)


def build_doc():
    ensure_dirs()
    create_limit_cycle_figure()
    doc = Document()
    style_doc(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("如何理解 Hopf 振荡器的极限环")
    r.bold = True
    set_east_asia(r)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("结合 SmartRollator 助行车 CPG 变导纳控制代码")
    sr.font.size = Pt(13)
    sr.font.color.rgb = RGBColor(75, 85, 99)
    set_east_asia(sr)

    add_callout(
        doc,
        "一句话定义",
        "极限环是一条稳定的闭合轨道。系统状态不收敛到某一个静止点，而是收敛到这条轨道上，并沿着它持续周期运动。对于 Hopf 振荡器，这条轨道通常就是 x-y 平面中的一个圆。",
    )

    add_p(
        doc,
        "在你的助行车项目里，极限环可以理解为 CPG 的“稳定节律来源”。它保证 x、y 不会发散，也不会停在一个固定值，而是形成连续、可预测、有频率的周期状态。随后代码只取 x 去调制导纳阻尼 D，使助行车的助力带有步态节律。"
    )

    doc.add_heading("1. 先用直观语言理解极限环", level=1)
    add_p(doc, "普通稳定点的含义是：系统最后停在一个点。例如小球滚进碗底，最后静止在最低点。极限环不一样，它的稳定目标不是一个点，而是一圈闭合轨道。")
    add_p(doc, "可以把它想象成一个带自动修正能力的跑道：")
    add_bullets(
        doc,
        [
            "状态在跑道里面时，系统把它往外推，直到回到跑道。",
            "状态在跑道外面时，系统把它往内拉，直到回到跑道。",
            "状态已经在跑道上时，它不会停下来，而是沿着跑道持续转动。",
        ],
    )
    add_p(doc, "因此，极限环的关键词是：有界、周期、抗扰动、可恢复。")

    doc.add_heading("2. Hopf 振荡器里的极限环是什么", level=1)
    add_p(doc, "你当前代码使用的 Hopf 形式是：")
    add_formula(
        doc,
        [
            "dx/dt = α(μ - x² - y²)x - ωy + εF_ext",
            "dy/dt = α(μ - x² - y²)y + ωx",
            "r² = x² + y²",
        ],
    )
    add_p(doc, "这里 x 和 y 不是两个独立输出，而是同一个二维振荡器的状态。把 x 当横轴、y 当纵轴，系统每一时刻的状态就是平面上的一个点。这个点最终会靠近一个圆，并沿这个圆持续旋转。这个圆就是 Hopf 振荡器的极限环。")
    add_formula(doc, ["稳定半径：r = sqrt(μ)", "默认 μ = 1.0，因此默认极限环半径约为 1"])

    doc.add_picture(str(FIG_PATH), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run("图 1：Hopf 极限环的直观解释及其在助行车代码中的映射")
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(75, 85, 99)
    set_east_asia(cr)

    doc.add_heading("3. 代码里极限环是怎么体现的", level=1)
    add_p(doc, "当前工程中，极限环主要由三段代码体现。")
    add_table(
        doc,
        ["代码逻辑", "对应数学意义", "为什么重要"],
        [
            ("r2 = x*x + y*y", "计算当前状态离原点的半径平方 r²。", "判断当前状态是在目标环内、环上，还是环外。"),
            ("alpha * (mu - r2)", "径向恢复项。", "r² < μ 时向外推；r² > μ 时向内拉；使系统回到极限环。"),
            ("-omega*y 和 +omega*x", "旋转项。", "让状态沿 x-y 平面持续转动，形成周期节律。"),
            ("x += dx*dt; y += dy*dt", "欧拉积分更新状态。", "每 1 ms 推进一次振荡器状态。"),
        ],
        [1.65, 2.05, 2.2],
    )
    add_p(doc, "这也是为什么代码必须保留 y。没有 y，系统就很难形成二维闭合轨道，只会更像一个向固定值收敛的一维系统。x 和 y 一起构成极限环，最后工程上只把 x 拿出来调制阻尼。")

    doc.add_heading("4. α、μ、ω 对极限环分别做什么", level=1)
    add_table(
        doc,
        ["参数", "对极限环的影响", "对助行车体验的影响"],
        [
            ("μ", "决定极限环半径，稳定半径 r = sqrt(μ)。", "μ 越大，x 的振幅越大，阻尼调制越强；太大会有忽轻忽重感。"),
            ("ω", "决定沿极限环转动的速度，不主要改变半径。", "决定助力节律快慢；应尽量与人体步频匹配。"),
            ("α", "决定偏离极限环后回到轨道的速度。", "α 小则恢复慢、跟随迟钝；α 大则恢复快，但可能放大噪声或抖动。"),
            ("x", "极限环上的横向状态。", "当前用于 D = D_base - λx，决定某一相位下车更轻还是更稳。"),
            ("y", "极限环上的纵向状态，保存相位。", "不直接输出到车轮，但保证 x 的变化连续、光滑、有相位记忆。"),
        ],
        [0.75, 2.45, 2.7],
    )
    add_p(doc, "可以把 μ、ω、α 的分工理解为：μ 管“跑道有多大”，ω 管“沿跑道跑多快”，α 管“跑偏后被拉回跑道有多快”。")

    doc.add_heading("5. 为什么极限环适合助行车", level=1)
    add_p(doc, "助行车需要的不是一个随意振荡的信号，而是一个稳定、可控、抗扰动的节律。极限环正好满足这个要求。")
    add_bullets(
        doc,
        [
            "稳定性：即使初始 x=0.1、y=0，状态也会逐渐靠近半径 sqrt(μ) 的轨道。",
            "有界性：状态不会无限增大，代码还额外加入了 CPG_STATE_LIMIT、CPG_DERIV_LIMIT 等防发散限幅。",
            "节律性：状态沿环持续旋转，因此 x 会形成周期变化。",
            "可调性：通过 μ 调幅值，通过 ω 调频率，通过 α 调恢复速度。",
            "可同步性：后续动捕步频输入只需要更新 ω，就能让节律跟随人体步频变化。",
        ],
    )

    doc.add_heading("6. 极限环如何影响车子的最终感觉", level=1)
    add_p(doc, "在当前代码中，极限环不会直接把车轮变成来回转。真正的车轮速度仍然来自导纳模型。极限环通过 x 改变的是阻尼：")
    add_formula(
        doc,
        [
            "D_current = D_base - λx",
            "M · dv_c/dt + D_current · v_c = F_ext",
            "v_c 再换算为左右轮 rpm",
        ],
    )
    add_table(
        doc,
        ["极限环状态", "阻尼变化", "用户感觉"],
        [
            ("x 为正", "D_current 下降", "车更轻、更容易被推走，助力感更明显。"),
            ("x 为负", "D_current 上升", "车更稳、更不容易突然窜动。"),
            ("ω 与步频匹配", "阻尼变化节奏贴近人的步伐", "推车感觉顺，不容易出现反拍感。"),
            ("α 合适", "扰动后能平稳回到节律", "经过传感器噪声或推力变化后，车不会长期乱掉。"),
            ("μ 或 λ 过大", "阻尼周期变化过强", "可能出现明显忽轻忽重，需要降低调制强度。"),
        ],
        [1.3, 1.7, 3.0],
    )

    doc.add_heading("7. 如果没有极限环，会发生什么", level=1)
    add_p(doc, "如果没有 Hopf 极限环，控制系统一般会变成两类情况：")
    add_bullets(
        doc,
        [
            "只用固定导纳：车可以根据力输出速度，但没有主动节律，不能自然匹配人的步态周期。",
            "直接用普通正弦信号：可以有节律，但缺少受扰后自动恢复到稳定轨道的非线性机制。",
        ],
    )
    add_p(doc, "Hopf 极限环的优势在于：它不是简单播放一个固定波形，而是一个动态系统。受到扰动后，它能通过自身方程回到稳定轨道，这比纯正弦发生器更适合人机交互场景。")

    doc.add_heading("8. 汇报时可以这样讲", level=1)
    add_p(
        doc,
        "极限环可以理解为 Hopf 振荡器在 x-y 状态平面中的稳定闭合轨道。系统不会停在一个平衡点，而是收敛到这条轨道并持续旋转。μ 决定轨道半径，ω 决定沿轨道运动的频率，α 决定偏离轨道后恢复的快慢。在我的助行车代码中，x 和 y 共同形成极限环，但最终只使用 x 去调制导纳阻尼 D。这样做的意义是让助行车获得稳定的步态节律调制，同时避免直接把周期信号当轮速导致车轮前后往复。"
    )

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("SmartRollator Hopf 极限环说明")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(100, 116, 139)
    set_east_asia(fr)

    doc.save(DOCX_PATH)
    print(DOCX_PATH)
    print(FIG_PATH)


if __name__ == "__main__":
    build_doc()
