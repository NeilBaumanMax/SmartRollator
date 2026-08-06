from __future__ import annotations

import math
import os
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "generated_docs"
ASSET_DIR = OUT_DIR / "assets"
DOCX_PATH = OUT_DIR / "Hopf振荡器参数详解_智能助行车.docx"
FIG_PATH = ASSET_DIR / "hopf_parameter_influence.png"


DEFAULTS = {
    "alpha": 10.0,
    "mu": 1.0,
    "omega": 2.0 * math.pi * 0.9,
    "epsilon": 0.01,
    "x": 0.1,
    "y": 0.0,
    "M": 15.0,
    "D_base": 40.0,
    "lambda": 20.0,
    "v_c": 0.0,
    "D_min": 15.0,
    "D_max": 80.0,
    "v_min": -0.5,
    "v_max": 1.2,
}


def ensure_dirs() -> None:
    OUT_DIR.mkdir(exist_ok=True)
    ASSET_DIR.mkdir(exist_ok=True)


def find_font() -> str:
    candidates = [
        "/System/Library/Fonts/STHeiti Medium.ttc",
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Supplemental/Arial.ttf",
    ]
    for path in candidates:
        if Path(path).exists():
            return path
    return ""


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    path = find_font()
    if path:
        return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def draw_center(draw: ImageDraw.ImageDraw, box, text, fill, fnt) -> None:
    x1, y1, x2, y2 = box
    bbox = draw.multiline_textbbox((0, 0), text, font=fnt, spacing=6)
    tw = bbox[2] - bbox[0]
    th = bbox[3] - bbox[1]
    draw.multiline_text(
        (x1 + (x2 - x1 - tw) / 2, y1 + (y2 - y1 - th) / 2),
        text,
        font=fnt,
        fill=fill,
        align="center",
        spacing=6,
    )


def draw_arrow(draw: ImageDraw.ImageDraw, p1, p2, fill=(70, 87, 105), width=4) -> None:
    draw.line([p1, p2], fill=fill, width=width)
    x1, y1 = p1
    x2, y2 = p2
    ang = math.atan2(y2 - y1, x2 - x1)
    size = 13
    pts = [
        (x2, y2),
        (x2 - size * math.cos(ang - 0.45), y2 - size * math.sin(ang - 0.45)),
        (x2 - size * math.cos(ang + 0.45), y2 - size * math.sin(ang + 0.45)),
    ]
    draw.polygon(pts, fill=fill)


def create_parameter_figure(path: Path) -> None:
    img = Image.new("RGB", (1800, 1120), "white")
    draw = ImageDraw.Draw(img)
    title = font(42)
    h = font(28)
    body = font(23)
    small = font(19)

    draw.text((70, 46), "Hopf 振荡器参数如何影响极限环与助行车体验", font=title, fill=(17, 24, 39))
    draw.text(
        (72, 102),
        "核心逻辑：x、y 形成稳定节律；μ 决定极限环半径；ω 决定沿环运动快慢；α 决定回到极限环的速度；x 进一步调制导纳阻尼 D。",
        font=body,
        fill=(75, 85, 99),
    )

    # Panel 1: phase plane and limit cycle
    panel1 = (70, 170, 840, 640)
    draw.rounded_rectangle(panel1, radius=18, outline=(37, 99, 235), width=3, fill=(248, 251, 255))
    draw.text((100, 198), "1. 极限环：μ 决定半径", font=h, fill=(30, 64, 175))
    cx, cy, scale = 455, 420, 155
    draw.line((cx - 230, cy, cx + 230, cy), fill=(148, 163, 184), width=2)
    draw.line((cx, cy + 190, cx, cy - 190), fill=(148, 163, 184), width=2)
    draw.text((cx + 238, cy - 8), "x", font=small, fill=(51, 65, 85))
    draw.text((cx + 9, cy - 218), "y", font=small, fill=(51, 65, 85))
    draw.ellipse((cx - scale, cy - scale, cx + scale, cy + scale), outline=(37, 99, 235), width=5)
    draw.text((cx - 110, cy + scale + 24), "r = sqrt(μ)", font=body, fill=(30, 64, 175))
    for mul, col, label in [(0.55, (251, 146, 60), "振幅小：向外"), (1.45, (239, 68, 68), "振幅大：向内")]:
        px = cx + int(scale * mul / math.sqrt(2))
        py = cy - int(scale * mul / math.sqrt(2))
        tx = cx + int(scale / math.sqrt(2))
        ty = cy - int(scale / math.sqrt(2))
        draw_arrow(draw, (px, py), (tx, ty), fill=col, width=3)
        draw.text((px + 10, py - 30), label, font=small, fill=col)
    draw.arc((cx - scale, cy - scale, cx + scale, cy + scale), start=300, end=30, fill=(16, 185, 129), width=7)
    draw_arrow(draw, (cx + 126, cy - 85), (cx + 145, cy - 30), fill=(16, 185, 129), width=3)

    # Panel 2: time waveform and omega
    panel2 = (960, 170, 1730, 640)
    draw.rounded_rectangle(panel2, radius=18, outline=(16, 185, 129), width=3, fill=(247, 254, 251))
    draw.text((990, 198), "2. 节律速度：ω 决定周期", font=h, fill=(4, 120, 87))
    x0, y0 = 1020, 440
    draw.line((x0, y0, x0 + 600, y0), fill=(148, 163, 184), width=2)
    draw.line((x0, y0 - 150, x0, y0 + 150), fill=(148, 163, 184), width=2)
    points = []
    for i in range(0, 600):
        t = i / 600 * 2 * math.pi * 2.4
        points.append((x0 + i, y0 - int(95 * math.sin(t))))
    draw.line(points, fill=(16, 185, 129), width=4)
    draw_arrow(draw, (x0 + 390, y0 - 118), (x0 + 500, y0 - 118), fill=(4, 120, 87), width=3)
    draw.text((x0 + 350, y0 - 158), "ω 变大：周期更短，节律更快", font=small, fill=(4, 120, 87))
    draw.text((x0 + 30, y0 + 105), "T = 2π / ω；代码默认 f = 0.9 Hz", font=body, fill=(51, 65, 85))

    # Panel 3: damping modulation
    panel3 = (70, 710, 840, 1040)
    draw.rounded_rectangle(panel3, radius=18, outline=(234, 88, 12), width=3, fill=(255, 251, 247))
    draw.text((100, 738), "3. 车端映射：x 调制导纳阻尼", font=h, fill=(194, 65, 12))
    boxes = [
        ((120, 820, 290, 920), "x > 0\nD 降低"),
        ((350, 820, 520, 920), "D = D_base\n- λx"),
        ((580, 820, 750, 920), "x < 0\nD 升高"),
    ]
    for box, label in boxes:
        draw.rounded_rectangle(box, radius=12, outline=(234, 88, 12), width=2, fill=(255, 255, 255))
        draw_center(draw, box, label, (124, 45, 18), body)
    draw_arrow(draw, (290, 870), (350, 870), fill=(234, 88, 12), width=3)
    draw_arrow(draw, (520, 870), (580, 870), fill=(234, 88, 12), width=3)
    draw.text((120, 962), "阻尼低：更轻、更跟手；阻尼高：更稳、更不容易被扰动拖走", font=body, fill=(75, 85, 99))

    # Panel 4: user experience chain
    panel4 = (960, 710, 1730, 1040)
    draw.rounded_rectangle(panel4, radius=18, outline=(124, 58, 237), width=3, fill=(250, 247, 255))
    draw.text((990, 738), "4. 参数最终影响使用体验", font=h, fill=(109, 40, 217))
    chain = [
        ((990, 820, 1138, 920), "把手力\nF_ext"),
        ((1185, 820, 1333, 920), "Hopf\n节律状态"),
        ((1380, 820, 1528, 920), "动态阻尼\nD"),
        ((1575, 820, 1700, 920), "轮速\nrpm"),
    ]
    for box, label in chain:
        draw.rounded_rectangle(box, radius=12, outline=(124, 58, 237), width=2, fill=(255, 255, 255))
        draw_center(draw, box, label, (76, 29, 149), body)
    for a, b in zip(chain[:-1], chain[1:]):
        draw_arrow(draw, (a[0][2], 870), (b[0][0], 870), fill=(124, 58, 237), width=3)
    draw.text((990, 962), "调得好：助力随步态变轻/变稳；调得差：拖拽感、忽快忽慢或响应迟钝。", font=body, fill=(75, 85, 99))

    img.save(path)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D7DEE8", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_east_asia(run, font_name="Microsoft YaHei") -> None:
    run.font.name = font_name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), font_name)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(17, 24, 39)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Title", 24, "0B2545", 0, 12),
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        st = styles[name]
        st.font.name = "Calibri"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def add_paragraph(doc, text: str = "", style=None, bold_prefix: str | None = None):
    p = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        set_east_asia(r1)
        r2 = p.add_run(text[len(bold_prefix):])
        set_east_asia(r2)
    else:
        r = p.add_run(text)
        set_east_asia(r)
    return p


def add_formula_block(doc, lines: list[str]) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    table.columns[0].width = Inches(6.3)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F7FB")
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    set_table_borders(table, color="CBD5E1", size="8")
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, line in enumerate(lines):
        if i:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Cambria Math"
        run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(15, 23, 42)
    doc.add_paragraph()


def add_callout(doc, title: str, body: str, fill="F4F7FB", border="93C5FD") -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    table.columns[0].width = Inches(6.3)
    set_table_borders(table, color=border, size="10")
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=160, bottom=160, start=180, end=180)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(30, 64, 175)
    set_east_asia(r)
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    r2.font.size = Pt(10.5)
    set_east_asia(r2)
    doc.add_paragraph()


def set_col_widths(table, widths_inch):
    for row in table.rows:
        for idx, width in enumerate(widths_inch):
            row.cells[idx].width = Inches(width)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.allow_autofit = False
    set_col_widths(table, widths)
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(header)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(30, 64, 175)
        set_east_asia(r)
    for row_data in rows:
        row = table.add_row()
        for i, value in enumerate(row_data):
            cell = row.cells[i]
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            if i == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(value)
            r.font.size = Pt(9.2)
            set_east_asia(r)
    doc.add_paragraph()
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(item)
        set_east_asia(r)


def build_doc() -> None:
    ensure_dirs()
    create_parameter_figure(FIG_PATH)

    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("Hopf 振荡器参数详解")
    r.bold = True
    set_east_asia(r)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("面向 STM32H7 智能助行车 CPG 变导纳控制")
    sr.font.size = Pt(13)
    sr.font.color.rgb = RGBColor(75, 85, 99)
    set_east_asia(sr)

    add_callout(
        doc,
        "汇报主线",
        "Hopf 振荡器在本项目中不是直接产生轮速，而是产生一个稳定、可调频率、可调幅值的节律状态。该节律状态通过 x 调制导纳阻尼 D，使助行车的助力响应能够随人体步态节奏变化。",
    )

    add_paragraph(
        doc,
        "本文解释图中 x、y、μ、ω、α 五个核心参数，并补充说明代码中为了把 Hopf 振荡器接入助行车而加入的 ε、D_base、λ、M、v_c 等参数。重点不是单纯背公式，而是说明每个参数为什么这样定义、它对极限环有什么影响，以及最后会怎样改变用户推车时的感觉。"
    )

    doc.add_heading("1. 本项目采用的 Hopf-CPG 结构", level=1)
    add_paragraph(doc, "标准 Hopf 振荡器由两个状态变量 x、y 构成，核心方程为：")
    add_formula_block(
        doc,
        [
            "dx/dt = α(μ - x² - y²)x - ωy + εF_ext",
            "dy/dt = α(μ - x² - y²)y + ωx",
            "r² = x² + y²,    稳态半径 r = sqrt(μ)",
        ],
    )
    add_paragraph(
        doc,
        "其中 εF_ext 是本工程加入的外力耦合项，用来让把手力传感器输入能够轻微影响振荡状态。没有外力输入时，系统会依靠自身动力学收敛到稳定极限环；有外力输入时，振荡器会在节律基础上被人体推力轻微牵引。"
    )
    add_paragraph(doc, "在助行车里，Hopf 状态进入导纳控制的方式是：")
    add_formula_block(
        doc,
        [
            "D_current = D_base - λx",
            "M · dv_c/dt + D_current · v_c = F_ext",
            "v_c → 轮速换算 → 左右轮 CAN 速度指令",
        ],
    )
    add_paragraph(
        doc,
        "因此，x 不是直接等于车轮速度。x 是节律调制信号，真正输出给车轮的是导纳模型积分得到的 v_c。这样做可以保留 Hopf 的周期节律，又避免车轮因为 x 的正负振荡而前后往复运动。"
    )

    doc.add_heading("2. 参数总览：数学含义、代码默认值、车端作用", level=1)
    rows = [
        ("x", "0.1", "无量纲状态", "主振荡状态", "用于调制阻尼 D，决定当前节律相位下助行车更轻还是更稳。"),
        ("y", "0", "无量纲状态", "辅助振荡状态", "和 x 共同形成二维极限环；当前未直接接车轮，但决定相位连续性。"),
        ("μ", "1.0", "状态幅值平方", "目标极限环半径平方", "决定 x、y 的稳定振幅；μ 越大，阻尼周期变化越明显。"),
        ("ω", "2π×0.9", "rad/s", "固有角频率", "决定节律快慢；后续接动捕步频时主要更新它。"),
        ("α", "10.0", "1/s 量级", "收敛速度", "决定系统偏离极限环后恢复速度；过小迟钝，过大易抖。"),
        ("ε", "0.01", "状态/力耦合", "外力耦合系数", "决定把手力对 Hopf 状态的牵引强度。"),
        ("D_base", "40.0", "N·s/m", "基础阻尼", "决定没有节律调制时导纳系统的基本厚重感。"),
        ("λ", "20.0", "阻尼调制增益", "x 到 D 的映射增益", "决定节律对阻尼的影响强弱。"),
        ("M", "15.0", "kg 等效量", "虚拟质量", "决定速度响应惯性；越大越平顺但越慢。"),
    ]
    add_table(doc, ["参数", "默认值", "单位/类型", "数学作用", "助行车作用"], rows, [0.65, 0.9, 1.05, 1.7, 2.0])

    doc.add_picture(str(FIG_PATH), width=Inches(6.3))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cr = cap.add_run("图 1：参数对极限环、阻尼调制和轮速输出链路的影响")
    cr.font.size = Pt(9)
    cr.font.color.rgb = RGBColor(75, 85, 99)
    set_east_asia(cr)

    doc.add_heading("3. x：主振荡状态", level=1)
    add_paragraph(
        doc,
        "x 是 Hopf 振荡器的主要输出状态。在传统 CPG 机器人控制中，x 常被直接映射为关节角度、关节速度或相位信号。但在助行车中，如果把 x 直接映射成轮速，车轮会随着正弦节律出现正转、反转交替，这与助行车连续前进的物理需求冲突。"
    )
    add_paragraph(
        doc,
        "所以本工程把 x 定义为“阻尼调制变量”，而不是“速度指令变量”。x 为正时，D_current = D_base - λx 会降低阻尼，用户会感觉车更容易被推动；x 为负时，阻尼升高，车更稳、更不容易被扰动带走。"
    )
    add_formula_block(doc, ["x > 0  →  D_current 降低  →  更轻、更跟手", "x < 0  →  D_current 升高  →  更稳、更保守"])
    add_bullets(
        doc,
        [
            "对振荡器的影响：x 是二维状态之一，和 y 一起决定当前相位和半径。",
            "对极限环的影响：x 本身不是极限环半径，x² + y² 才决定是否在极限环上。",
            "对使用体验的影响：x 的正负变化会让助力呈现周期性的轻/稳切换，而不是让轮子周期性倒转。",
        ],
    )

    doc.add_heading("4. y：辅助状态变量", level=1)
    add_paragraph(
        doc,
        "y 是 x 的配套状态变量。它的意义不是“另一个传感器量”，而是让系统形成二维旋转动力学。如果只有 x，一个一维系统通常只能趋向某个平衡点，很难自然形成稳定的周期轨道；引入 y 后，状态点可以在 x-y 平面中绕原点旋转，并被径向项拉回固定半径。"
    )
    add_formula_block(doc, ["相位可理解为：φ = atan2(y, x)", "半径可理解为：r = sqrt(x² + y²)"])
    add_bullets(
        doc,
        [
            "对振荡器的影响：y 与 x 相互耦合，-ωy 和 +ωx 两项共同产生旋转。",
            "对极限环的影响：y 参与半径 r 的计算，决定系统是否偏离目标圆环。",
            "对使用体验的影响：虽然当前代码没有直接使用 y，但 y 保证 x 的变化是连续、光滑、有相位记忆的，而不是突变的开关信号。",
        ],
    )

    doc.add_heading("5. μ：目标振幅平方", level=1)
    add_paragraph(
        doc,
        "μ 是 Hopf 振荡器最关键的幅值参数。它被写在 μ - x² - y² 这一项里，原因是 x² + y² 正好表示当前半径平方 r²。这样系统就可以根据“当前半径平方”和“目标半径平方”的差值自动调整振幅。"
    )
    add_formula_block(doc, ["dr/dt = α(μ - r²)r", "稳定时：dr/dt = 0  →  r = sqrt(μ)"])
    add_paragraph(
        doc,
        "这也是为什么图中说最终振幅是 sqrt(μ)，而不是 μ 本身。把 μ 定义成半径平方，可以让方程写得更简洁，也让幅值控制和频率控制天然解耦。"
    )
    add_table(
        doc,
        ["μ 取值", "极限环表现", "车端表现"],
        [
            ("μ < 0", "原点稳定，不产生持续振荡。", "没有周期性阻尼调制，不适合 CPG 助力。"),
            ("μ = 0", "临界情况，振幅趋近 0。", "节律基本消失，接近普通固定导纳。"),
            ("μ > 0", "产生稳定极限环，半径 sqrt(μ)。", "存在稳定节律调制，适合步态同步助力。"),
        ],
        [1.0, 2.45, 2.45],
    )
    add_bullets(
        doc,
        [
            "μ 太小：x 的振幅小，D_current 变化弱，用户几乎感觉不到 CPG 的节律作用。",
            "μ 适中：阻尼有轻微周期变化，助行车既能跟手，又不会明显忽快忽慢。",
            "μ 太大：x 的振幅过大，阻尼调制过强，车可能出现明显的周期性轻重变化，使用体验会变得突兀。",
        ],
    )

    doc.add_heading("6. ω：固有角频率", level=1)
    add_paragraph(
        doc,
        "ω 决定状态点沿极限环旋转的快慢，也就是节律频率。它不主要改变极限环的大小，而是改变同一个圆环被走完一圈需要多长时间。"
    )
    add_formula_block(doc, ["ω = 2πf", "T = 2π / ω = 1 / f"])
    add_paragraph(
        doc,
        f"当前代码默认 f = 0.9 Hz，因此 ω = 2π×0.9 ≈ {DEFAULTS['omega']:.3f} rad/s。这相当于让助力节律以接近慢走步频的速度运行。"
    )
    add_bullets(
        doc,
        [
            "对振荡器的影响：ω 越大，x 和 y 的周期越短，节律越快。",
            "对极限环的影响：理想情况下，ω 改变沿环运动速度，不改变极限环半径。",
            "对使用体验的影响：ω 与人体步频匹配时，用户会感觉助力节律顺着步伐；ω 偏低会像车反应慢半拍，ω 偏高会像车在主动催促或轻微拉扯。",
        ],
    )
    add_paragraph(
        doc,
        "这就是为什么后续动捕系统接入时，最应该实时更新的是 ω，而不是 μ。更新 ω 可以让节律跟随人的步频变化，同时不破坏振幅稳定性。"
    )

    doc.add_heading("7. α：收敛速度", level=1)
    add_paragraph(
        doc,
        "α 控制系统偏离极限环后回到稳定圆环的快慢。它只出现在径向修正项 α(μ - x² - y²) 中，因此主要影响“恢复速度”，不直接决定最终振幅，也不直接决定频率。"
    )
    add_formula_block(doc, ["当前半径太小：r² < μ  →  α(μ-r²) 为正  →  振幅增大", "当前半径太大：r² > μ  →  α(μ-r²) 为负  →  振幅减小"])
    add_table(
        doc,
        ["α 取值倾向", "振荡器表现", "车端体验"],
        [
            ("偏小", "回到极限环慢，对外力扰动恢复慢。", "助力节律跟随慢，感觉迟钝或拖泥带水。"),
            ("适中", "能较快恢复，但不会造成明显数值抖动。", "节律稳定，推车感觉连续自然。"),
            ("偏大", "恢复很快，但欧拉积分下更容易受噪声和步长影响。", "可能出现抖动、突变或助力忽然变轻/变重。"),
        ],
        [1.15, 2.4, 2.35],
    )
    add_paragraph(
        doc,
        "当前代码采用 1 ms 控制周期和一阶欧拉积分，α = 10 是一个偏稳妥的初始值。如果后续传感器噪声明显，α 不宜盲目增大，否则 Hopf 状态会更积极地响应扰动，反而放大体感上的不平顺。"
    )

    doc.add_heading("8. 为什么说 Hopf 的幅值和频率是解耦的", level=1)
    add_paragraph(
        doc,
        "从极坐标角度看，Hopf 振荡器可以近似理解为两个独立问题：半径 r 怎么回到目标值，以及相位 φ 怎么匀速转动。μ 和 α 主要管半径，ω 主要管相位速度。"
    )
    add_formula_block(doc, ["dr/dt = α(μ-r²)r", "dφ/dt = ω"])
    add_paragraph(
        doc,
        "这就是它适合做人机步态节律控制的原因：当动捕系统测到人的步频变化时，只需要更新 ω，就可以改变节律快慢；而 μ 不变时，节律幅值仍保持稳定。反过来，如果只想改变阻尼调制强度，可以调 μ 或 λ，而不必改变步频。"
    )

    doc.add_heading("9. 对极限环的集中影响", level=1)
    add_table(
        doc,
        ["参数", "是否改变极限环半径", "是否改变绕行速度", "是否改变恢复速度", "说明"],
        [
            ("x", "否", "间接体现相位", "否", "x 是当前点的位置，不是极限环设计参数。"),
            ("y", "否", "间接体现相位", "否", "y 与 x 共同定义当前相位和半径。"),
            ("μ", "是", "否", "间接影响", "稳定半径 r = sqrt(μ)。"),
            ("ω", "否", "是", "否", "改变相位速度和周期 T。"),
            ("α", "否", "否", "是", "改变偏离极限环后的收敛快慢。"),
            ("εF_ext", "会造成扰动/牵引", "可能牵引相位", "取决于输入", "外力输入会让轨迹相对标准圆环发生轻微变形。"),
        ],
        [0.75, 1.25, 1.15, 1.15, 1.9],
    )

    doc.add_heading("10. 对助行车最终使用体验的影响", level=1)
    add_paragraph(
        doc,
        "用户真正感受到的不是 x、y 或极限环本身，而是轮子输出速度和阻尼变化。参数调节最终会通过 D_current 和 v_c 体现为“推起来轻不轻、跟不跟手、稳不稳、是否有拖拽感”。"
    )
    add_table(
        doc,
        ["体验问题", "可能相关参数", "原因", "调整方向"],
        [
            ("车太沉、推不动", "D_base 偏大，λ 偏小，μ 偏小", "阻尼整体高或节律调制弱。", "适当降低 D_base，或小幅增大 λ/μ。"),
            ("车忽轻忽重", "μ 或 λ 偏大", "x 对阻尼调制过强。", "优先减小 λ，再考虑减小 μ。"),
            ("节律慢半拍", "ω 偏低，α 偏小", "节律频率或恢复速度跟不上人体步态。", "用动捕步频更新 ω，必要时小幅增加 α。"),
            ("车有抖动感", "α 偏大，ε 偏大，力传感器噪声大", "系统过度响应扰动或数值变化太快。", "降低 α/ε，对力信号滤波。"),
            ("有前后往复趋势", "v_min 允许负速度，或误把 x 映射成轮速", "x 本身正负振荡，不能直接当轮速。", "助行车若只允许前进，应把速度下限改为 0。"),
        ],
        [1.1, 1.5, 1.95, 1.45],
    )
    add_callout(
        doc,
        "重要工程判断",
        "如果你的真实助行车不允许倒车，建议把代码中的 CPG_V_MIN 从 -0.5f 改为 0.0f。这样 Hopf 仍然可以调制阻尼，但导纳输出速度不会给出反向车速。",
        fill="FFF7ED",
        border="FDBA74",
    )

    doc.add_heading("11. 建议调参顺序", level=1)
    add_paragraph(doc, "为了避免多个参数同时变化导致现象难以解释，建议按下面顺序调试：")
    steps = [
        "先固定 F_ext = 0，观察 x-y 是否能从初值收敛到半径约 sqrt(μ) 的极限环。",
        "固定 μ = 1、α = 10，只调 ω，确认节律频率能跟预期步频一致。",
        "接入力传感器，但先保持 ε 较小，避免外力直接把极限环拉得过度变形。",
        "调 D_base，让车在无明显节律调制时有合适的基础厚重感。",
        "再调 λ，让 x 对阻尼的调制达到“能感觉到但不突兀”的程度。",
        "最后根据实车体验微调 μ 和 α；μ 管节律幅度，α 管恢复快慢。",
    ]
    for idx, step in enumerate(steps, start=1):
        p = doc.add_paragraph(style="List Number")
        r = p.add_run(step)
        set_east_asia(r)

    doc.add_heading("12. 和当前代码的对应关系", level=1)
    add_paragraph(doc, "当前工程中的核心实现位于：MDK-ARM/cpg_control.c 和 MDK-ARM/cpg_control.h。关键对应关系如下：")
    add_table(
        doc,
        ["代码位置/变量", "作用", "说明"],
        [
            ("CPG_DEFAULT_ALPHA", "设置 α = 10.0", "控制收敛速度。"),
            ("CPG_DEFAULT_MU", "设置 μ = 1.0", "目标极限环半径为 1。"),
            ("CPG_DEFAULT_OMEGA", "设置 ω = 2π×0.9", "默认节律频率为 0.9 Hz。"),
            ("CPG_Step()", "每 1 ms 更新一次 Hopf 和导纳", "先更新 x/y，再计算 D_current，最后积分得到 v_c。"),
            ("CPG_Update_Frequency()", "外部步频接入口", "接收到 0.3 到 2.0 Hz 的人体步频后更新 ω。"),
            ("D_current = D_base - λx", "CPG 到导纳的桥梁", "这是本项目把 Hopf 用到助行车上的核心映射。"),
        ],
        [1.6, 1.85, 2.45],
    )

    doc.add_heading("13. 汇报时可以这样总结", level=1)
    add_paragraph(
        doc,
        "Hopf 振荡器提供的是一个稳定极限环节律。μ 决定这个节律的幅值，ω 决定节律频率，α 决定受到扰动后回到稳定节律的快慢。x、y 是振荡器内部状态，其中 x 被我映射到导纳阻尼，使助行车在人的步态节律中周期性改变柔顺程度。这样既保留了 CPG 的稳定节律，又避免了直接用振荡信号驱动车轮造成的前后往复问题。最终，用户感受到的是更贴合步态的助力，而不是单纯固定阻尼的机械响应。"
    )

    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("SmartRollator Hopf-CPG 参数说明")
    fr.font.size = Pt(9)
    fr.font.color.rgb = RGBColor(100, 116, 139)
    set_east_asia(fr)

    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
    print(DOCX_PATH)
    print(FIG_PATH)
