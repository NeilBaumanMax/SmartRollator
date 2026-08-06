from pathlib import Path
from datetime import datetime
import math

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


ROOT = Path("/Users/neil/项目/SmartRollator/SmartRollator")
DOC_DIR = Path("/Users/neil/Documents")
OUT = DOC_DIR / "智能助行器Hopf-CPG变导纳控制实验设计与数据处理方案.docx"
DIAGRAM_DIR = ROOT / "diagrams"
FLOW_IMG = DIAGRAM_DIR / "experiment_design_flow.png"
MATRIX_IMG = DIAGRAM_DIR / "experiment_matrix_overview.png"


def set_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_para_spacing(paragraph, before=0, after=6, line=1.16):
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line


def add_p(doc, text="", style=None, size=10.5, bold=False, color=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_font(r, size=size, bold=bold, color=color)
    return p


def add_formula(doc, formula, note=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    set_para_spacing(p, after=3)
    r = p.add_run(formula)
    set_font(r, name="Cambria Math", size=11, bold=False, color="1F4D78")
    if note:
        pn = doc.add_paragraph()
        pn.paragraph_format.left_indent = Inches(0.18)
        set_para_spacing(pn, after=6)
        rn = pn.add_run(note)
        set_font(rn, size=9.5, color="57606A")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_para_spacing(p, after=3)
    r = p.add_run(text)
    set_font(r, size=10)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_para_spacing(p, after=3)
    r = p.add_run(text)
    set_font(r, size=10)
    return p


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margins(cell, top=90, start=120, bottom=90, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for k, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{k}"))
        if node is None:
            node = OxmlElement(f"w:{k}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def table_borders(table, color="D0D7DE", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_width(table, widths_in):
    for row in table.rows:
        for cell, width in zip(row.cells, widths_in):
            cell.width = Inches(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def add_table(doc, rows, widths, header=True, font_size=9):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_borders(table)
    table.autofit = False
    for i, val in enumerate(rows[0]):
        c = table.rows[0].cells[i]
        c.text = ""
        if header:
            shade_cell(c, "F2F4F7")
        cell_margins(c)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = c.paragraphs[0]
        set_para_spacing(p, after=0, line=1.05)
        r = p.add_run(str(val))
        set_font(r, size=font_size, bold=header, color="0B2545" if header else None)
    for row in rows[1:]:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            c = cells[i]
            c.text = ""
            cell_margins(c)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = c.paragraphs[0]
            set_para_spacing(p, after=0, line=1.05)
            r = p.add_run(str(val))
            set_font(r, size=font_size)
    set_table_width(table, widths)
    doc.add_paragraph()
    return table


def style_doc(doc):
    sec = doc.sections[0]
    sec.top_margin = Inches(0.9)
    sec.bottom_margin = Inches(0.85)
    sec.left_margin = Inches(0.85)
    sec.right_margin = Inches(0.85)
    sec.header_distance = Inches(0.49)
    sec.footer_distance = Inches(0.49)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(6)

    style_specs = [
        ("Title", 22, "0B2545", 0, 8, True),
        ("Subtitle", 11, "57606A", 0, 10, False),
        ("Heading 1", 16, "2E74B5", 14, 7, True),
        ("Heading 2", 13, "2E74B5", 10, 5, True),
        ("Heading 3", 11.5, "1F4D78", 8, 4, True),
    ]
    for name, size, color, before, after, bold in style_specs:
        st = styles[name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.font.bold = bold
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def draw_rounded_rect(draw, box, radius, fill, outline, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def get_fonts():
    candidates = [
        "/System/Library/Fonts/PingFang.ttc",
        "/System/Library/Fonts/STHeiti Light.ttc",
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
    ]
    for p in candidates:
        try:
            return (
                ImageFont.truetype(p, 42),
                ImageFont.truetype(p, 28),
                ImageFont.truetype(p, 22),
                ImageFont.truetype(p, 18),
            )
        except OSError:
            pass
    f = ImageFont.load_default()
    return f, f, f, f


def draw_images():
    DIAGRAM_DIR.mkdir(exist_ok=True)
    title, head, body, small = get_fonts()

    # Flow image.
    img = Image.new("RGB", (1800, 820), "white")
    d = ImageDraw.Draw(img)
    d.text((70, 45), "论文实验主线：从算法验证到实车人机实验", fill=(11, 37, 69), font=title)
    steps = [
        ("E0", "离线仿真", "验证 Hopf 极限环、参数影响、数值稳定性"),
        ("E1", "台架验证", "轮子悬空，检查方向、同步、CAN 下发和速度响应"),
        ("E2", "传感器标定", "力传感器零偏、线性、重复性；视觉/动捕时间同步"),
        ("E3", "直线推行", "固定导纳 vs Hopf-CPG，对比力、速度、节律跟随"),
        ("E4", "变步频推行", "慢速/舒适/快速，验证 omega 接入和相位适应"),
        ("E5", "主观评价", "稳定性、轻便性、自然性、疲劳感量表"),
        ("D", "数据处理", "滤波、分段、指标计算、统计检验、论文图表"),
    ]
    x0, y0 = 75, 180
    box_w, box_h, gap = 220, 155, 30
    colors = [(232, 239, 245), (245, 241, 232), (238, 245, 239), (242, 238, 248), (248, 238, 238), (237, 245, 245), (244, 246, 248)]
    for i, (code, name, desc) in enumerate(steps):
        x = x0 + i * (box_w + gap)
        if x + box_w > 1725:
            x = x0 + (i - 4) * (box_w + gap)
            y = y0 + 260
        else:
            y = y0
        draw_rounded_rect(d, (x, y, x + box_w, y + box_h), 18, colors[i], (112, 132, 155), 2)
        d.text((x + 22, y + 16), code, fill=(31, 77, 120), font=head)
        d.text((x + 22, y + 57), name, fill=(36, 41, 47), font=head)
        # simple wrap
        chars = 13
        lines = [desc[j:j + chars] for j in range(0, len(desc), chars)]
        for k, line in enumerate(lines[:3]):
            d.text((x + 22, y + 100 + 26 * k), line, fill=(87, 96, 106), font=small)
        if i < len(steps) - 1:
            if i == 3:
                d.line((x + box_w // 2, y + box_h + 10, x + box_w // 2, y + box_h + 95), fill=(112, 132, 155), width=5)
                d.polygon([(x + box_w // 2, y + box_h + 105), (x + box_w // 2 - 12, y + box_h + 83), (x + box_w // 2 + 12, y + box_h + 83)], fill=(112, 132, 155))
            elif i != 6:
                nx = x + box_w + gap
                if nx + box_w <= 1725:
                    d.line((x + box_w + 5, y + box_h // 2, nx - 5, y + box_h // 2), fill=(112, 132, 155), width=5)
                    d.polygon([(nx - 5, y + box_h // 2), (nx - 25, y + box_h // 2 - 12), (nx - 25, y + box_h // 2 + 12)], fill=(112, 132, 155))
    d.text((70, 745), "原则：先证明算法和硬件可靠，再做人体推行；每个实验都保留固定导纳对照组，数据用同一套指标处理。", fill=(87, 96, 106), font=body)
    img.save(FLOW_IMG)

    # Matrix image.
    img2 = Image.new("RGB", (1800, 960), "white")
    d2 = ImageDraw.Draw(img2)
    d2.text((70, 45), "实验变量矩阵与主要评价指标", fill=(11, 37, 69), font=title)
    # axes
    d2.text((120, 150), "控制策略", fill=(31, 77, 120), font=head)
    d2.text((560, 150), "步行任务", fill=(31, 77, 120), font=head)
    d2.text((1000, 150), "采集数据", fill=(31, 77, 120), font=head)
    d2.text((1390, 150), "评价指标", fill=(31, 77, 120), font=head)
    cols = [
        (90, 200, 410, 660, ["无助力/低助力", "固定导纳", "Hopf-CPG 固定频率", "Hopf-CPG 步频同步"]),
        (530, 200, 850, 660, ["静态推力", "台架阶跃", "10 m 直线", "慢/舒适/快步频", "起停与转弯"]),
        (970, 200, 1290, 660, ["F_x/F_y/F_z", "x/y/r/omega/D", "v_cmd/rpm_cmd", "rpm_fb/速度", "视觉/动捕/视频"]),
        (1370, 200, 1690, 660, ["推力均值/RMS/峰值", "速度平稳性", "同步误差", "相位差/节律跟随", "主观自然性/稳定性"]),
    ]
    for x1, y1, x2, y2, items in cols:
        draw_rounded_rect(d2, (x1, y1, x2, y2), 20, (244, 246, 248), (208, 215, 222), 2)
        for j, item in enumerate(items):
            yy = y1 + 35 + j * 78
            draw_rounded_rect(d2, (x1 + 28, yy, x2 - 28, yy + 50), 12, (255, 255, 255), (208, 215, 222), 1)
            d2.text((x1 + 48, yy + 12), item, fill=(36, 41, 47), font=body)
    d2.line((420, 430, 520, 430), fill=(112, 132, 155), width=5)
    d2.polygon([(520, 430), (500, 418), (500, 442)], fill=(112, 132, 155))
    d2.line((860, 430, 960, 430), fill=(112, 132, 155), width=5)
    d2.polygon([(960, 430), (940, 418), (940, 442)], fill=(112, 132, 155))
    d2.line((1300, 430, 1360, 430), fill=(112, 132, 155), width=5)
    d2.polygon([(1360, 430), (1340, 418), (1340, 442)], fill=(112, 132, 155))
    d2.text((100, 735), "推荐论文主假设：与固定导纳相比，Hopf-CPG 变导纳能降低推行力峰值和速度波动，并提高步态节律匹配与主观自然性。", fill=(31, 77, 120), font=head)
    d2.text((100, 790), "关键控制：同一受试者内随机化顺序；每种条件至少 3 次重复；每次记录完整时间戳和原始数据；先做安全限速和急停。", fill=(87, 96, 106), font=body)
    img2.save(MATRIX_IMG)


def build_doc():
    DOC_DIR.mkdir(parents=True, exist_ok=True)
    draw_images()

    doc = Document()
    style_doc(doc)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("智能助行器 Hopf-CPG 变导纳控制")
    set_font(r, size=22, bold=True, color="0B2545")
    title2 = doc.add_paragraph(style="Title")
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title2.add_run("实验设计与数据处理方案")
    set_font(r, size=21, bold=True, color="0B2545")

    sub = doc.add_paragraph(style="Subtitle")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("面向论文实验章节、实车测试 SOP 与数据分析流程")
    set_font(r, size=11, color="57606A")

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(f"生成日期：{datetime.now().strftime('%Y-%m-%d')}    工程：SmartRollator / STM32H7 / ZLAC8015D / USART2-RS485 力传感器 / USART1 视觉输入")
    set_font(r, size=9.5, color="57606A")

    add_p(doc, "使用说明", style="Heading 1")
    add_p(doc, "本文档按论文可发表实验的逻辑组织：先验证算法本身，再验证硬件执行，再进行人体推行对比实验，最后用统一的数据处理方法得到统计指标。方案默认你的当前工程已经实现：右把手力传感器 F_x 输入、Hopf-CPG 变导纳控制、1 ms 控制周期、CANopen 速度模式下发左右轮 rpm、USART1 周期日志和视觉位姿接收。")

    add_p(doc, "一、研究目标与论文假设", style="Heading 1")
    add_p(doc, "本研究的核心问题是：在智能助行器中，引入基于 Hopf 振荡器的 CPG 变导纳控制后，是否能在保持安全稳定的前提下，使助行器的助力更符合人体步态节律，并降低用户推行负担。")
    add_bullet(doc, "目标 1：证明 Hopf 振荡器在 STM32H7 离散控制周期下能够形成稳定极限环，且 mu、omega、alpha 对振幅、频率和收敛速度的影响符合理论。")
    add_bullet(doc, "目标 2：证明 CPG 状态 x 调制导纳阻尼后，助行器在相同外力输入下具有更符合步态节律的速度响应。")
    add_bullet(doc, "目标 3：与固定导纳控制相比，验证 Hopf-CPG 变导纳控制是否降低推行力峰值、推行力 RMS、速度波动和主观疲劳感。")
    add_bullet(doc, "目标 4：验证后续接入动捕步频后，通过 CPG_Update_Frequency() 更新 omega 能提高控制节律与人体步频的一致性。")
    add_p(doc, "建议论文假设：H1：Hopf-CPG 变导纳控制相较固定导纳控制能显著降低用户推行力 RMS；H2：Hopf-CPG 变导纳控制能降低速度波动并提升速度平滑性；H3：步频同步的 Hopf-CPG 控制比固定频率 Hopf-CPG 具有更小的人机节律相位差。")

    add_p(doc, "二、当前工程基础与实验变量", style="Heading 1")
    add_table(doc, [
        ["模块", "当前工程状态", "实验中需要记录/确认的内容"],
        ["力传感器", "USART2 RS485 查询，约 100 Hz，当前控制使用右把手 F_x。", "记录 fx、fy、fz、timestamp；实验前做零偏和标定。"],
        ["控制循环", "ControlTask 1 ms 周期，调用 Controller_Run()。", "记录每次控制的 F_x、v_cmd、rpm_cmd、控制模式。"],
        ["Hopf-CPG", "cpg_control.c 中更新 x、y、omega、D_current、v_c。", "论文实验建议额外输出 x、y、r、omega、D_current、v_c。"],
        ["电机驱动", "FDCAN1/CANopen 速度模式，左右轮目标 rpm 幅值一致，右轮符号反向补偿安装方向。", "记录左/右目标 rpm，若能读反馈则记录 rpm_fb_L/R。"],
        ["视觉输入", "USART1 接收 ARUCO 位姿与 dist_mm，目前主要用于日志快照。", "可作为位移/距离辅助验证；主实验可用外部视频或动捕作为地面真值。"],
        ["上位机记录", "USART1 周期日志默认 100 ms。", "正式实验建议改为 10 ms 或 20 ms CSV 输出，避免动态指标采样不足。"],
    ], widths=[1.2, 2.95, 3.45])

    add_p(doc, "控制策略自变量", style="Heading 2")
    add_table(doc, [
        ["编号", "控制策略", "说明", "论文用途"],
        ["C0", "无助力/低助力安全模式", "电机不主动输出或仅保持极低辅助，用户手推。", "作为人体自然推行基线，非必须但推荐。"],
        ["C1", "固定导纳控制", "M = 15.0，D = 40.0，不使用 CPG 调制。", "核心对照组。"],
        ["C2", "Hopf-CPG 固定频率变导纳", "alpha=10，mu=1，omega=2π×0.9，D=D_base−lambda*x。", "验证 CPG 变导纳相对固定导纳的效果。"],
        ["C3", "Hopf-CPG 步频同步变导纳", "外部动捕/步频估计输入 freq_hz，调用 CPG_Update_Frequency() 更新 omega。", "验证人体节律同步能力，若动捕暂未接入可列为扩展实验。"],
    ], widths=[0.65, 1.8, 3.0, 2.15])

    add_p(doc, "因变量与评价指标", style="Heading 2")
    add_table(doc, [
        ["类别", "指标", "计算意义"],
        ["推行负担", "F_x 均值、F_x RMS、F_x 峰值、单位距离力积分", "衡量用户推动助行器所需力量。"],
        ["运动性能", "平均速度、速度标准差、加速度 RMS、jerk RMS", "衡量车体运动是否平稳。"],
        ["控制性能", "v_cmd、rpm_cmd、rpm_fb、速度跟踪误差、左右轮同步误差", "衡量控制输出和执行器一致性。"],
        ["CPG 特性", "r=sqrt(x²+y²)、r 稳态误差、omega、D_current、相位差", "证明 Hopf 极限环和步频同步是否有效。"],
        ["安全性", "最大速度、急停次数、异常限幅次数、传感器超时次数", "证明算法不会产生危险输出。"],
        ["主观体验", "轻便性、稳定性、自然性、可控性、疲劳感 Likert 评分", "补充说明用户使用体验。"],
    ], widths=[1.1, 2.4, 4.1])

    if FLOW_IMG.exists():
        add_p(doc, "实验总流程图", style="Heading 2")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(FLOW_IMG), width=Inches(7.0))

    add_p(doc, "三、实验环境与设备准备", style="Heading 1")
    add_p(doc, "硬件准备", style="Heading 2")
    for item in [
        "STM32H7 开发板、DM-MC-Board02 电机开发板、ZLAC8015D 驱动器、左右轮电机、力传感器、RS485 接口、CAN 接口、供电电池。保留机械急停或软件急停。 ",
        "上位机：用于记录 USART1 日志，推荐使用 Python 串口脚本或串口助手保存原始 CSV。串口参数按当前工程 USART1：115200、8N1。 ",
        "地面真值设备：优先使用动作捕捉系统；没有动捕时，使用固定相机视频 + 地面 1 m 标尺/ArUco 标记；若视觉测距稳定，也可记录 dist_mm 辅助验证。 ",
        "实验场地：至少 12 m 平直走廊或实验室通道，取中间 10 m 为有效测试段，两端各留 1 m 加速/减速缓冲。地面保持平整、干燥，无障碍物。 ",
        "安全人员：人体实验时至少 1 人在旁保护，第一次上电测试必须轮子悬空。 ",
    ]:
        add_bullet(doc, item)

    add_p(doc, "软件准备", style="Heading 2")
    for item in [
        "固件版本锁定：记录实验当天 git/工程版本或至少保存 MDK-ARM/cpg_control.c、Controller.c、Core/Src/freertos.c 的副本。 ",
        "日志频率调整：正式动态实验建议 LOG_PERIOD_MS 从 100 ms 改为 10 ms 或 20 ms；如果担心串口带宽，至少记录 50 Hz。 ",
        "CSV 输出建议字段：t_ms, trial_id, condition, fx, fy, fz, x_cpg, y_cpg, r_cpg, omega, D_current, v_cmd_mps, rpm_cmd_L, rpm_cmd_R, rpm_fb_L, rpm_fb_R, vision_dist_mm, mocap_freq_hz, event_flag。 ",
        "CPG 调试字段：当前 cpg_control.c 内 D_current 是局部变量，正式实验建议增加 CPG_Debug_t 或 getter，把 x、y、omega、D_current、v_c 输出给日志。 ",
        "安全限幅确认：速度上限 v_c <= 1.2 m/s，电机 rpm 限幅 400 rpm；如果论文实验只允许前进，建议把 CPG_V_MIN 改为 0.0f，避免反向速度。 ",
    ]:
        add_bullet(doc, item)

    add_p(doc, "四、实验前标定与安全检查 SOP", style="Heading 1")
    add_p(doc, "每次实验开始前按以下顺序执行，任何一项不通过都不进入人体实验。")
    pre_steps = [
        "机械检查：确认轮胎固定、线缆固定、手柄无松动、电池电量充足、急停按钮有效。",
        "轮子悬空上电：将助行器架起，使两个驱动轮离地；打开电源，烧录或确认固件版本。",
        "CAN 方向检查：下发 +30 rpm，观察左右轮物理前进方向是否一致；如果右轮命令为负但物理前进，记录为安装方向补偿，不视为反转问题。",
        "力传感器零偏采集：手柄无外力保持 10 s，记录 fx、fy、fz，计算零偏 offset_Fx、offset_Fy、offset_Fz；正式数据处理中统一扣除。",
        "力传感器线性检查：用已知砝码或拉力计施加 5 N、10 N、20 N 三个等级，每级保持 5 s，检查输出均值与理论值的线性误差。",
        "速度安全检查：轮子悬空情况下分别运行 C1、C2，确认 rpm_cmd 不超过限幅，左右轮目标 rpm 幅值一致。",
        "日志检查：发送 p 命令或开启流式日志，确认每行包含时间戳、力、命令；正式实验前保存 30 s 空载日志作为噪声基线。",
        "场地检查：标记起点、终点和有效 10 m 区间；地面无杂物；受试者穿防滑鞋。",
        "人体试推：受试者以最低速度试推 2 m，确认手柄、急停、保护人员位置和数据记录均正常。",
    ]
    for s in pre_steps:
        add_number(doc, s)

    add_p(doc, "五、实验 E0：Hopf 振荡器离线仿真验证", style="Heading 1")
    add_p(doc, "目的：在不接入硬件的条件下证明算法符合 Hopf 振荡器理论，避免把实车问题和算法问题混在一起。")
    add_p(doc, "操作步骤", style="Heading 2")
    for s in [
        "用 MATLAB/Python/C 复现 CPG_Step() 中的欧拉积分，dt 设置为 0.001 s，总仿真时间 20 s。",
        "输入 F_ext = 0，参数使用默认值 alpha=10，mu=1，omega=2π×0.9，epsilon=0.01，初始 x=0.1，y=0。",
        "记录 x、y、r=sqrt(x²+y²)、theta=atan2(y,x)、D_current、v_c。",
        "改变 mu：0.25、1.0、2.25，验证稳定半径分别约为 0.5、1.0、1.5。",
        "改变 omega：2π×0.6、2π×0.9、2π×1.2，验证振荡频率变化但半径基本不变。",
        "改变 alpha：3、10、20，比较 r 收敛到 sqrt(mu) 所需时间。",
        "加入外力扰动 F_ext：在 5 s 到 6 s 加入 20 N 阶跃，观察系统偏离极限环后能否恢复。",
    ]:
        add_number(doc, s)
    add_p(doc, "输出图表：x-y 相图、r(t) 曲线、x(t) 曲线、不同 mu/omega/alpha 的对比曲线。")

    add_p(doc, "六、实验 E1：台架电机与控制链路验证", style="Heading 1")
    add_p(doc, "目的：验证 STM32H7 控制循环、CANopen 下发、左右轮同步和安全限幅，不涉及人体风险。")
    add_p(doc, "操作步骤", style="Heading 2")
    for s in [
        "将助行器架起，保证驱动轮完全离地，连接上位机串口并开始保存日志。",
        "运行固定导纳 C1。用手沿前进方向缓慢推动右把手力传感器，持续 5 s 后松开，重复 5 次。",
        "运行 Hopf-CPG C2。重复同样的 5 次推力输入。",
        "如果具备电机反馈读取，记录 rpm_fb_L 与 rpm_fb_R；没有反馈时，用高速视频或光电测速仪记录轮速作为备选。",
        "分别测试小推力、中等推力、较大推力三档，目标是覆盖 5 N、15 N、30 N 左右的手柄力范围。",
        "每次测试后检查日志是否连续，有无传感器超时、CAN 错误、rpm 限幅触发。",
    ]:
        add_number(doc, s)
    add_p(doc, "判定标准：左右轮目标 rpm 幅值误差应为 0；若有反馈，稳态左右轮反馈速度差建议小于 5%。任何条件下速度命令不得超过预设安全上限。")

    add_p(doc, "七、实验 E2：传感器标定与同步验证", style="Heading 1")
    add_p(doc, "目的：保证后续人体实验数据可信。传感器标定不是论文亮点，但如果不做，推力和速度指标的可信度会下降。")
    add_table(doc, [
        ["项目", "操作", "输出结果"],
        ["力传感器零偏", "无外力静置 10 s，采样至少 1000 点。", "offset_Fx、offset_Fy、offset_Fz；噪声标准差 sigma_F。"],
        ["力传感器线性", "用拉力计或砝码施加 5/10/20/30 N，每档 5 s，重复 3 次。", "线性拟合斜率、截距、R²、最大误差。"],
        ["视觉/动捕同步", "实验开始前做一次明显推车前进-停止事件，同时记录 STM32 与动捕/视频时间。", "时间偏移 Δt_sync，用于后处理对齐。"],
        ["轮速/车速校准", "固定 rpm 运行，测量 3 m 行驶时间或用动捕测速度。", "轮半径修正值 R_wheel、rpm 到 m/s 的换算误差。"],
    ], widths=[1.35, 3.45, 2.8])

    add_p(doc, "八、实验 E3：10 m 直线推行对比实验", style="Heading 1")
    add_p(doc, "这是论文的主实验。采用同一受试者内对照设计，比较固定导纳和 Hopf-CPG 变导纳控制的推行负担与运动平稳性。")
    add_p(doc, "受试者建议", style="Heading 2")
    add_bullet(doc, "预实验：3 到 5 名健康成年人，用于检查流程、限幅和指标。")
    add_bullet(doc, "正式实验：建议 12 到 20 名健康成年人；如果有真实目标人群，需先通过伦理审批并降低速度上限。")
    add_bullet(doc, "纳入条件：无严重下肢损伤，能独立完成 10 m 步行。排除条件：近期跌倒风险、眩晕、严重心肺疾病。")
    add_p(doc, "实验条件", style="Heading 2")
    add_table(doc, [
        ["条件", "控制策略", "任务速度", "重复次数"],
        ["C1-S", "固定导纳", "慢速，约 0.5 m/s", "3 次"],
        ["C1-N", "固定导纳", "舒适速度，约 0.8 m/s", "3 次"],
        ["C1-F", "固定导纳", "快速但安全，约 1.0 m/s", "3 次"],
        ["C2-S", "Hopf-CPG 固定频率", "慢速", "3 次"],
        ["C2-N", "Hopf-CPG 固定频率", "舒适速度", "3 次"],
        ["C2-F", "Hopf-CPG 固定频率", "快速但安全", "3 次"],
        ["C3-N", "Hopf-CPG 步频同步", "舒适速度", "3 次，可作为扩展条件"],
    ], widths=[1.0, 2.2, 2.65, 1.35])
    add_p(doc, "操作步骤", style="Heading 2")
    for s in [
        "给受试者讲解任务：自然握住手柄，按指定速度推行 10 m，不刻意配合算法。",
        "进行 2 次熟悉试走，不记录为正式数据。",
        "按随机顺序执行各控制条件，避免学习效应。可以把条件编号写在纸条上随机抽取。",
        "每次 trial 开始前，受试者站在起点前 1 m；上位机创建数据文件，命名如 S01_C2_N_T03.csv。",
        "记录员口令“开始”，受试者推行通过起点，进入有效 10 m 区间；到达终点后继续走出 1 m 缓冲区再停止。",
        "每次 trial 后立即保存日志，并记录异常事件：急停、脚步打乱、旁人干预、传感器掉线。",
        "同一条件的 3 次重复之间休息 30 到 60 s，不同控制策略之间休息 2 min，避免疲劳影响。",
        "每完成一个控制策略，让受试者填写 1 到 7 分主观评分：轻便性、稳定性、自然性、可控性、疲劳感。",
    ]:
        add_number(doc, s)

    add_p(doc, "九、实验 E4：变步频与 CPG 同步实验", style="Heading 1")
    add_p(doc, "目的：验证 omega = 2πf 的设计是否能让 CPG 节律跟随人体步频。该实验依赖动捕、足底开关、IMU 或视频算法提取步频。如果暂时没有动捕，可先用视频离线提取步频，作为论文扩展实验。")
    for s in [
        "受试者在 10 m 通道内按节拍器行走，节拍设置为 0.6 Hz、0.9 Hz、1.2 Hz 三档。",
        "每档先让受试者原地或空推适应 20 s，再开始正式试验。",
        "条件 A：Hopf-CPG 固定 omega = 2π×0.9，不接入步频。",
        "条件 B：实时输入 freq_hz，调用 CPG_Update_Frequency()，使 omega = 2π×freq_hz。",
        "每档步频、每种条件重复 3 次，顺序随机化。",
        "记录 CPG 相位 theta_cpg=atan2(y,x) 和人体步态相位 theta_gait，后处理计算相位差。",
    ]:
        add_number(doc, s)
    add_p(doc, "判定标准：与固定 omega 相比，步频同步条件下 CPG 频率误差和相位差应更小，且推行力和速度波动不应变差。")

    add_p(doc, "十、实验 E5：起停、转弯与安全性补充实验", style="Heading 1")
    add_p(doc, "这部分不一定作为论文主结果，但建议用于说明系统安全性和工程完整性。")
    add_table(doc, [
        ["实验", "操作", "关注指标"],
        ["起步", "从静止开始推行 3 m，重复 5 次。", "起步响应时间、最大加速度、推力峰值。"],
        ["停止", "正常推行后停止施力，观察车速衰减。", "停止时间、残余速度、是否继续前冲。"],
        ["转弯", "绕 1.5 m 半径圆弧行走或 90 度转弯。", "左右轮速度一致策略是否需要差速扩展。"],
        ["异常力", "短时快速推拉手柄，保护人员在旁。", "限幅是否触发，v_c 是否保持在安全范围。"],
        ["传感器丢包", "临时断开或屏蔽力传感器输入，仅台架测试。", "freshness 判断是否清零输出。"],
    ], widths=[1.1, 3.6, 2.9])

    if MATRIX_IMG.exists():
        add_p(doc, "实验变量矩阵", style="Heading 2")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(MATRIX_IMG), width=Inches(7.0))

    add_p(doc, "十一、数据记录格式与字段定义", style="Heading 1")
    add_p(doc, "正式实验建议全部保存为 CSV。每一行对应一个采样时刻，所有设备尽量统一到 STM32 时间戳或后处理对齐后的时间轴。")
    add_table(doc, [
        ["字段", "单位", "来源", "用途"],
        ["t_ms", "ms", "STM32 HAL_GetTick 或上位机时间", "统一时间轴。"],
        ["trial_id", "-", "上位机", "区分受试者、条件和重复。"],
        ["condition", "-", "上位机/固件", "C0/C1/C2/C3。"],
        ["fx, fy, fz", "N", "USART2 力传感器", "推行负担指标。"],
        ["x_cpg, y_cpg", "-", "CPG_Handle_t", "极限环和相位分析。"],
        ["r_cpg", "-", "sqrt(x²+y²)", "极限环半径稳定性。"],
        ["omega", "rad/s", "CPG_Handle_t", "步频同步分析。"],
        ["D_current", "N·s/m", "CPG_Step 内部计算", "变导纳机理解释。"],
        ["v_cmd_mps", "m/s", "CPG 输出", "速度命令分析。"],
        ["rpm_cmd_L/R", "rpm", "MotorTxTask", "电机命令一致性。"],
        ["rpm_fb_L/R", "rpm", "驱动器反馈，可选", "速度跟踪和同步误差。"],
        ["vision_dist_mm", "mm", "USART1 ARUCO", "位移/距离辅助验证。"],
        ["mocap_freq_hz", "Hz", "动捕或视频", "步频输入和人机同步分析。"],
        ["event_flag", "-", "人工标记", "起点、终点、异常、急停。"],
    ], widths=[1.45, 0.85, 2.15, 3.15], font_size=8.6)

    add_p(doc, "CSV 文件命名规则", style="Heading 2")
    add_formula(doc, "S{受试者编号}_{控制条件}_{速度任务}_T{重复次数}.csv", "例：S01_C2_N_T03.csv 表示 1 号受试者，Hopf-CPG 固定频率，舒适速度，第 3 次重复。")

    add_p(doc, "十二、数据预处理方法", style="Heading 1")
    add_p(doc, "1. 数据导入与完整性检查", style="Heading 2")
    for s in [
        "检查每个 CSV 是否包含必要字段，删除空文件和明显中断文件。",
        "根据 t_ms 计算采样间隔 dt_i，统计丢包率。若连续丢包超过 200 ms，该 trial 标记为无效。",
        "检查 fx、v_cmd、rpm_cmd 是否存在 NaN、inf 或超出物理范围的值；异常值先标记，不直接删除。",
    ]:
        add_number(doc, s)
    add_p(doc, "2. 时间同步", style="Heading 2")
    add_formula(doc, "t_aligned = t_device + Δt_sync", "Δt_sync 由实验开始前的明显同步事件估计，例如推车启动瞬间、拍手声、LED 闪烁或动捕事件。")
    add_p(doc, "如果只有 STM32 日志，无外部动捕，则直接使用 t_ms。若有动捕或视频，使用互相关对齐：将车体速度或位移变化与视觉/动捕速度曲线做 cross-correlation，取相关峰值对应的时间偏移。")

    add_p(doc, "3. 零偏扣除与滤波", style="Heading 2")
    add_formula(doc, "F_x_corr(t) = F_x_raw(t) − mean(F_x_raw during zero-load)", "零偏来自每次实验前 10 s 无外力静置数据。")
    add_p(doc, "滤波建议：力信号使用 2 阶或 4 阶低通 Butterworth，截止频率 5 Hz；速度信号截止频率 3 到 5 Hz。若采样率较低，避免过高阶滤波造成相位畸变。论文中应说明滤波器阶数、截止频率和采样率。")

    add_p(doc, "4. 有效区间分段", style="Heading 2")
    for s in [
        "优先使用人工 event_flag 标记起点和终点。",
        "如果没有 event_flag，用位移达到 0.5 m 到 9.5 m 的区间作为有效段，排除起步和停止。",
        "如果只有速度日志，用 v_cmd 或估计车速超过 0.1 m/s 的时刻作为起点，速度连续低于 0.1 m/s 超过 0.5 s 作为终点。",
        "所有统计指标只在有效段内计算，起停实验除外。",
    ]:
        add_number(doc, s)

    add_p(doc, "十三、核心指标计算方法", style="Heading 1")
    add_p(doc, "推行负担指标", style="Heading 2")
    add_formula(doc, "F_mean = mean(|F_x_corr|)", "平均绝对推行力。")
    add_formula(doc, "F_RMS = sqrt(mean(F_x_corr²))", "推行力均方根，推荐作为主要推行负担指标。")
    add_formula(doc, "F_peak = max(|F_x_corr|)", "最大推行力，用于评价瞬时用力和安全性。")
    add_formula(doc, "W_F = ∫ |F_x_corr(t)| dt / distance", "单位距离推力积分，反映完成同样距离的总体用力。")

    add_p(doc, "运动平稳性指标", style="Heading 2")
    add_formula(doc, "v_mean = mean(v)", "平均行驶速度。v 可来自动捕/视觉，若没有外部测量，可先用 v_cmd 近似，但论文需说明。")
    add_formula(doc, "CV_v = std(v) / mean(v)", "速度变异系数，越小表示速度越稳定。")
    add_formula(doc, "a(t) = dv/dt,    J(t) = da/dt", "加速度和 jerk，可用差分后再滤波，或对平滑后的速度求导。")
    add_formula(doc, "J_RMS = sqrt(mean(J(t)²))", "jerk RMS 越小，推行越平顺。")

    add_p(doc, "控制性能指标", style="Heading 2")
    add_formula(doc, "rpm_cmd = v_cmd × 60 / (2πR)", "当前工程中 R = 0.075 m；若实测轮半径不同，应使用实测值。")
    add_formula(doc, "e_sync_cmd = ||rpm_cmd_L| − |rpm_cmd_R||", "左右轮命令同步误差；当前目标应接近 0。")
    add_formula(doc, "e_rpm = rpm_cmd − rpm_fb", "有电机反馈时计算速度跟踪误差。")
    add_formula(doc, "T_response = t(v reaches 90% steady value) − t(force step)", "台架阶跃输入下的响应时间。")

    add_p(doc, "CPG 与极限环指标", style="Heading 2")
    add_formula(doc, "r(t) = sqrt(x(t)² + y(t)²)", "CPG 当前半径。")
    add_formula(doc, "e_r = |r(t) − sqrt(mu)|", "极限环半径误差。")
    add_formula(doc, "theta_cpg(t) = atan2(y(t), x(t))", "CPG 相位。")
    add_formula(doc, "f_cpg = omega / (2π)", "CPG 当前频率。")
    add_formula(doc, "Δtheta = wrap(theta_cpg − theta_gait)", "有动捕或步态相位时，用于评价人机节律同步。")
    add_formula(doc, "D_current = D_base − lambda·x", "动态阻尼，用于解释为什么同样推力下速度响应会随步态节律变化。")

    add_p(doc, "主观评价处理", style="Heading 2")
    add_p(doc, "每个控制条件后填写 1 到 7 分量表：轻便性、稳定性、自然性、可控性、疲劳感。疲劳感是负向指标，统计总分时需要反向处理。也可增加 NASA-TLX 简化版，但论文工作量会增加。")

    add_p(doc, "十四、统计分析方法", style="Heading 1")
    for s in [
        "每名受试者、每个条件、每个速度任务下，先对 3 次重复取均值，避免单个受试者重复次数过多影响统计权重。",
        "主比较：C1 固定导纳 vs C2 Hopf-CPG 固定频率；扩展比较：C2 vs C3 步频同步。",
        "正态性检查：对差值进行 Shapiro-Wilk 检验。如果满足正态，用配对 t 检验；不满足正态，用 Wilcoxon 符号秩检验。",
        "多条件比较：若同时比较 C0/C1/C2/C3，使用重复测量 ANOVA；不满足正态或样本量小，使用 Friedman 检验。",
        "多重比较校正：使用 Bonferroni 或 Holm 校正。",
        "显著性水平：alpha = 0.05。报告 p 值时同时报告效应量和 95% 置信区间。",
        "效应量：配对 t 检验使用 Cohen's dz = mean(diff)/std(diff)；Wilcoxon 可报告 r = Z/sqrt(N)。",
        "论文图表建议统一报告 mean ± SD；如果样本较少，同时画出每名受试者的配对连线图。",
    ]:
        add_bullet(doc, s)

    add_p(doc, "十五、论文结果图表建议", style="Heading 1")
    add_table(doc, [
        ["图/表", "内容", "对应结论"],
        ["图 1", "系统结构和控制框图：力传感器 → Hopf-CPG → 动态阻尼 → 导纳速度 → CAN 电机。", "说明方法整体逻辑。"],
        ["图 2", "Hopf x-y 相图与 r(t) 收敛曲线。", "证明极限环稳定。"],
        ["图 3", "不同 mu、omega、alpha 的参数影响曲线。", "证明参数设计符合理论。"],
        ["图 4", "固定导纳和 Hopf-CPG 下 F_x、D_current、v_cmd 的时间序列。", "解释变导纳如何改变响应。"],
        ["图 5", "推行力 RMS、峰值、速度 CV 的箱线图或配对散点图。", "展示主要实验结果。"],
        ["图 6", "步频同步实验的 f_gait 与 f_cpg 曲线、相位差曲线。", "证明 omega 接入有效。"],
        ["表 1", "硬件参数与控制参数。", "保证实验可复现。"],
        ["表 2", "各控制条件下主要指标 mean ± SD 与统计检验结果。", "支撑论文结论。"],
        ["表 3", "主观评分结果。", "补充使用体验。"],
    ], widths=[1.0, 4.0, 2.6])

    add_p(doc, "十六、风险控制与无效数据判定", style="Heading 1")
    add_p(doc, "安全风险控制", style="Heading 2")
    for s in [
        "人体实验前必须完成台架验证，确认限速、急停、左右轮方向和传感器超时保护。",
        "第一次人体试验速度上限建议不超过 0.5 m/s；流程稳定后再提升到目标速度。",
        "受试者不适、步态明显不稳、设备异常响动、串口/传感器掉线、轮速异常时立即停止。",
        "所有测试旁边应有保护人员，急停按钮由记录员或保护人员掌握。",
    ]:
        add_bullet(doc, s)
    add_p(doc, "无效 trial 判定", style="Heading 2")
    for s in [
        "有效段内力传感器连续丢包超过 200 ms。",
        "有效段内速度命令触发安全限幅超过总时长 20%，说明该 trial 不是正常控制响应。",
        "受试者明显停步、转身、被保护人员扶持，或偏离测试路线。",
        "日志文件缺少关键字段，无法计算主要指标。",
        "视觉/动捕数据丢失时，该 trial 仍可用于力和命令指标，但不能用于步频同步指标。",
    ]:
        add_bullet(doc, s)

    add_p(doc, "十七、建议的实验实施顺序", style="Heading 1")
    for s in [
        "第 1 天：完成 E0 离线仿真和 E1 台架验证，修正日志字段和速度限幅。",
        "第 2 天：完成 E2 标定流程，得到力传感器零偏、线性系数和轮速换算误差。",
        "第 3 天：用 1 到 2 名同学做预实验，检查 10 m 推行流程和数据处理脚本。",
        "第 4 至 7 天：正式采集 12 到 20 名受试者的 E3 主实验数据。",
        "第 8 天：如果动捕/视频步频链路稳定，补采 E4 步频同步实验。",
        "第 9 天：完成数据清洗、指标计算和统计分析，输出论文图表。",
    ]:
        add_number(doc, s)

    add_p(doc, "十八、论文实验章节可直接采用的写法", style="Heading 1")
    add_p(doc, "本文采用受试者内对照实验设计，比较固定导纳控制与 Hopf-CPG 变导纳控制在智能助行器推行任务中的性能差异。实验系统由 STM32H7 控制器、RS485 力传感器、CANopen 电机驱动器和双轮差速底盘组成。控制周期为 1 ms，力传感器数据用于计算导纳速度命令，随后换算为左右轮目标转速。Hopf-CPG 控制器通过二维状态 x、y 产生稳定节律，并以 D_current = D_base − lambda·x 动态调制导纳阻尼。")
    add_p(doc, "每名受试者分别在固定导纳和 Hopf-CPG 变导纳条件下完成 10 m 直线推行任务。每种条件重复 3 次，实验顺序随机化。实验过程中同步记录手柄力、CPG 状态、动态阻尼、速度命令、电机转速命令和外部位移/步频数据。主要评价指标包括推行力 RMS、推行力峰值、速度变异系数、jerk RMS、左右轮同步误差和 CPG 极限环半径误差。统计分析采用配对检验或重复测量方差分析，显著性水平设为 0.05。")

    add_p(doc, "附录 A：建议优先补充的固件日志字段", style="Heading 1")
    add_table(doc, [
        ["优先级", "字段", "原因"],
        ["P0", "x_cpg, y_cpg", "没有这两个字段无法证明 Hopf 极限环实际运行。"],
        ["P0", "D_current", "没有该字段无法解释变导纳控制的核心作用。"],
        ["P0", "v_cmd_mps", "target_vel 只有 rpm，不利于和人体速度直接比较。"],
        ["P1", "omega/f_cpg", "步频同步实验必须记录。"],
        ["P1", "rpm_fb_L/R", "用于证明电机执行跟踪和左右轮同步。"],
        ["P2", "limit_flag/stale_flag", "用于安全性和无效数据筛查。"],
    ], widths=[0.9, 2.0, 4.7])

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    rr = footer.add_run("SmartRollator Hopf-CPG 实验设计与数据处理方案")
    set_font(rr, size=8.5, color="57606A")

    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    out = build_doc()
    print(out)
