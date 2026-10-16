from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


FILES_ROOT = Path(__file__).resolve().parents[1]
ROOT = FILES_ROOT.parent
HANDOVER = ROOT / "docs" / "handover" / "PROJECT_HANDOVER.md"
PROMPT = ROOT / "docs" / "handover" / "NEXT_AI_AGENT_PROMPT.md"
OUTPUT = FILES_ROOT / "handover" / "SmartRollator项目交接与Git重建指南.docx"

PAGE_W = 12240
PAGE_H = 15840
CONTENT_W = 9360
TABLE_INDENT = 120
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"


def rgb(value):
    return RGBColor.from_string(value)


def set_run_font(run, size=None, bold=None, italic=None, color=None, mono=False):
    latin = "Consolas" if mono else "Arial Unicode MS"
    east_asia = "Arial Unicode MS"
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), latin)
    rfonts.set(qn("w:hAnsi"), latin)
    rfonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = rgb(color)


def set_cell_margins(cell, top=80, bottom=80, start=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths):
    assert sum(widths) == CONTENT_W
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_W))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                cell._tc.get_or_add_tcPr().append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color="CBD5E1", size="5"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def add_numbering(doc, num_id, bullet=False):
    numbering = doc.part.numbering_part.element
    abstract_id = str(100 + num_id)
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), abstract_id)
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    lvl.append(start)
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if bullet else "decimal")
    lvl.append(num_fmt)
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if bullet else "%1.")
    lvl.append(lvl_text)
    suffix = OxmlElement("w:suff")
    suffix.set(qn("w:val"), "tab")
    lvl.append(suffix)
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    ppr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "540")
    ind.set(qn("w:hanging"), "270")
    ppr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.append(spacing)
    lvl.append(ppr)
    abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), abstract_id)
    num.append(abstract_ref)
    numbering.append(num)


def apply_num(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = ppr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        ppr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(number)


def add_inline_runs(paragraph, text, *, size=11, color=None, bold=False):
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=9.5, color="7C2D12", mono=True)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size, bold=bold, color=color)


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in tokens.items():
        style = styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial Unicode MS")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    add_numbering(doc, 21, bullet=True)
    add_numbering(doc, 22, bullet=False)


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    run = paragraph.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def set_running_header_footer(section):
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("SMARTROLLATOR  /  项目交接与 Git 重建")
    set_run_font(hr, size=8.5, bold=True, color=MUTED)
    hp.paragraph_format.space_after = Pt(0)
    fp = section.footer.paragraphs[0]
    add_page_field(fp)


def add_cover(doc):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("工程交接指南")
    set_run_font(r, size=11, bold=True, color=BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("SmartRollator")
    set_run_font(r, size=30, bold=True, color=INK)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(28)
    r = p.add_run("从零重建、Git 管理与 AI Agent 接手说明")
    set_run_font(r, size=15, color=DARK_BLUE)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("参考工程：STM32H723 + FreeRTOS + RS485 + CANopen + Hopf-CPG")
    set_run_font(r, size=10.5, color=MUTED)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(72)
    r = p.add_run("整理日期：2026-08-06  |  状态：可执行交接基线")
    set_run_font(r, size=10, color=MUTED)

    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = ""
    set_table_geometry(table, [CONTENT_W])
    set_table_borders(table, color="B8C7DA", size="7")
    shade_cell(table.cell(0, 0), CALLOUT)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_runs(p, "交接原则：旧工程只读；新仓库从空白 CubeMX 工程开始；每阶段独立提交并附验证证据。", size=11, color=INK, bold=True)
    doc.add_page_break()


def add_code_block(doc, lines):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_W])
    set_table_borders(table, color="D0D5DD", size="4")
    shade_cell(table.cell(0, 0), "F8FAFC")
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    for idx, line in enumerate(lines):
        if idx:
            p.add_run("\n")
        r = p.add_run(line)
        set_run_font(r, size=8.5, color="344054", mono=True)


def parse_table(lines, start):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        parts = [part.strip() for part in lines[i].strip().strip("|").split("|")]
        rows.append(parts)
        i += 1
    if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", x) for x in rows[1]):
        rows.pop(1)
    return rows, i


def column_widths(count):
    if count == 1:
        return [CONTENT_W]
    if count == 2:
        return [2700, 6660]
    if count == 3:
        return [1700, 2500, 5160]
    if count == 4:
        return [1250, 2150, 2980, 2980]
    base = CONTENT_W // count
    values = [base] * count
    values[-1] += CONTENT_W - sum(values)
    return values


def add_markdown_table(doc, rows):
    if not rows:
        return
    count = max(len(row) for row in rows)
    rows = [row + [""] * (count - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=count)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    set_table_geometry(table, column_widths(count))
    set_table_borders(table)
    for r_idx, values in enumerate(rows):
        for c_idx, value in enumerate(values):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            if r_idx == 0:
                shade_cell(cell, LIGHT_BLUE)
            elif r_idx % 2 == 0:
                shade_cell(cell, "FAFBFC")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.12
            add_inline_runs(p, value, size=9, bold=(r_idx == 0), color=INK if r_idx == 0 else None)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(1)


def add_blockquote(doc, text):
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [CONTENT_W])
    set_table_borders(table, color="9FB3CA", size="6")
    shade_cell(table.cell(0, 0), CALLOUT)
    p = table.cell(0, 0).paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    add_inline_runs(p, text, size=10.5, color=INK, bold=True)


def append_markdown(doc, path, skip_title=True):
    lines = path.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines = []
    first_heading = True
    current_decimal_num_id = None
    if not hasattr(doc, "_next_decimal_num_id"):
        doc._next_decimal_num_id = 30
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if line.startswith("```"):
            current_decimal_num_id = None
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if not line.strip() or line.strip() == "---":
            current_decimal_num_id = None
            i += 1
            continue
        if line.strip().startswith("|"):
            current_decimal_num_id = None
            rows, i = parse_table(lines, i)
            add_markdown_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            current_decimal_num_id = None
            raw_level = len(heading.group(1))
            level = max(1, raw_level - 1) if skip_title else raw_level
            title = heading.group(2)
            if first_heading and skip_title:
                first_heading = False
                i += 1
                continue
            first_heading = False
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline_runs(p, title, size={1: 16, 2: 13, 3: 12}[level], color=BLUE if level < 3 else DARK_BLUE, bold=True)
            i += 1
            continue
        if line.startswith("> "):
            current_decimal_num_id = None
            add_blockquote(doc, line[2:].strip())
            i += 1
            continue
        if re.match(r"^\s*- \[[ xX]\] ", line):
            current_decimal_num_id = None
            text = re.sub(r"^\s*- \[([ xX])\] ", lambda m: "☒ " if m.group(1).lower() == "x" else "☐ ", line)
            p = doc.add_paragraph()
            apply_num(p, 21)
            add_inline_runs(p, text, size=10.5)
            i += 1
            continue
        if re.match(r"^\s*- ", line):
            current_decimal_num_id = None
            text = re.sub(r"^\s*- ", "", line)
            p = doc.add_paragraph()
            apply_num(p, 21)
            add_inline_runs(p, text, size=10.5)
            i += 1
            continue
        if re.match(r"^\s*\d+\. ", line):
            text = re.sub(r"^\s*\d+\. ", "", line)
            p = doc.add_paragraph()
            if current_decimal_num_id is None:
                current_decimal_num_id = doc._next_decimal_num_id
                doc._next_decimal_num_id += 1
                add_numbering(doc, current_decimal_num_id, bullet=False)
            apply_num(p, current_decimal_num_id)
            add_inline_runs(p, text, size=10.5)
            i += 1
            continue
        p = doc.add_paragraph()
        current_decimal_num_id = None
        add_inline_runs(p, line.replace("  ", " "), size=11)
        i += 1


def build():
    doc = Document()
    configure_document(doc)
    set_running_header_footer(doc.sections[0])
    add_cover(doc)
    append_markdown(doc, HANDOVER, skip_title=True)
    doc.add_page_break()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("附录：给下一个 AI Agent 的启动提示词")
    set_run_font(r, size=22, bold=True, color=INK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("可整段复制到新任务中使用")
    set_run_font(r, size=11, color=MUTED)
    doc.add_page_break()
    append_markdown(doc, PROMPT, skip_title=True)
    core = doc.core_properties
    core.title = "SmartRollator 项目交接与 Git 重建指南"
    core.subject = "从零重建、Git 管理与 AI Agent 接手说明"
    core.author = "SmartRollator Project"
    core.keywords = "STM32H723, FreeRTOS, Hopf-CPG, Git, handover"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
