from pathlib import Path
from math import pi, cos, sin, sqrt

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "Hopf振荡器参数详解_助行车应用版.docx"
DIAGRAM_DIR = ROOT / "diagrams"
LIMIT_CYCLE_IMG = DIAGRAM_DIR / "hopf_parameter_limit_cycle.png"
LOGIC_IMG = DIAGRAM_DIR / "hopf_cpg_admittance_logic.png"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D0D7DE"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_run_font(run, name="Microsoft YaHei", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def add_para(doc, text="", style=None, size=None, bold=None, color=None, align=None):
    p = doc.add_paragraph(style=style)
    if text:
        r = p.add_run(text)
        set_run_font(r, size=size, bold=bold, color=color)
    if align is not None:
        p.alignment = align
    return p


def add_formula(doc, formula, note=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.left_indent = Inches(0.18)
    r = p.add_run(formula)
    set_run_font(r, name="Cambria Math", size=11, color="1F4D78")
    if note:
        pn = doc.add_paragraph()
        pn.paragraph_format.left_indent = Inches(0.18)
        pn.paragraph_format.space_after = Pt(6)
        rn = pn.add_run(note)
        set_run_font(rn, size=10, color="57606A")


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(text)
    set_run_font(r, size=10.5)
    return p


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(6)

    for style_name, size, color, before, after in [
        ("Title", 22, "0B2545", 0, 12),
        ("Subtitle", 11, "57606A", 0, 10),
        ("Heading 1", 16, "1F4D78", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 11.5, "1F4D78", 8, 4),
    ]:
        st = styles[style_name]
        st.font.name = "Microsoft YaHei"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor.from_string(color)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)


def draw_limit_cycle():
    DIAGRAM_DIR.mkdir(exist_ok=True)
    img = Image.new("RGB", (1800, 950), "white")
    d = ImageDraw.Draw(img)

    try:
        font_title = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 42)
        font_head = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 28)
        font_body = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 24)
        font_small = ImageFont.truetype("/System/Library/Fonts/PingFang.ttc", 20)
    except OSError:
        font_title = font_head = font_body = font_small = ImageFont.load_default()

    d.text((70, 45), "Hopf 振荡器参数、极限环与助行车体验的关系", fill=(11, 37, 69), font=font_title)

    # Phase plane.
    cx, cy = 470, 500
    scale = 190
    d.line((cx - 300, cy, cx + 300, cy), fill=(99, 110, 123), width=3)
    d.line((cx, cy + 300, cx, cy - 300), fill=(99, 110, 123), width=3)
    d.polygon([(cx + 305, cy), (cx + 285, cy - 10), (cx + 285, cy + 10)], fill=(99, 110, 123))
    d.polygon([(cx, cy - 305), (cx - 10, cy - 285), (cx + 10, cy - 285)], fill=(99, 110, 123))
    d.text((cx + 318, cy - 18), "x", fill=(87, 96, 106), font=font_body)
    d.text((cx + 12, cy - 342), "y", fill=(87, 96, 106), font=font_body)

    # Limit cycle r=sqrt(mu).
    r = scale
    d.ellipse((cx - r, cy - r, cx + r, cy + r), outline=(31, 77, 120), width=6)
    d.text((cx - 155, cy + r + 28), "极限环半径 r = sqrt(mu)", fill=(31, 77, 120), font=font_body)

    # Inner/outside arrows.
    for rr, color, text, pos in [
        (90, (36, 138, 61), "r < sqrt(mu)：振幅被推大", (120, 235)),
        (275, (191, 80, 35), "r > sqrt(mu)：振幅被拉回", (120, 735)),
    ]:
        for a in (0.2, 2.6, 4.7):
            x1 = cx + rr * cos(a)
            y1 = cy + rr * sin(a)
            x2 = cx + (rr + (25 if rr < r else -25)) * cos(a)
            y2 = cy + (rr + (25 if rr < r else -25)) * sin(a)
            d.line((x1, y1, x2, y2), fill=color, width=5)
            d.ellipse((x2 - 5, y2 - 5, x2 + 5, y2 + 5), fill=color)
        d.text(pos, text, fill=color, font=font_body)

    # Rotation arrows omega.
    pts = []
    for i in range(16):
        a = 5.2 + i * 0.035
        pts.append((cx + (r + 35) * cos(a), cy + (r + 35) * sin(a)))
    d.line(pts, fill=(88, 96, 196), width=5)
    d.polygon([(cx + 130, cy - 185), (cx + 158, cy - 180), (cx + 140, cy - 160)], fill=(88, 96, 196))
    d.text((cx + 105, cy - 260), "omega 决定旋转快慢", fill=(88, 96, 196), font=font_body)

    # Parameter cards.
    cards = [
        ("mu", "决定目标振幅：r = sqrt(mu)", "变大：极限环半径变大，x 的调制幅度增强", "对车：阻尼波动更明显，助力节律更强"),
        ("omega", "决定相位推进速度：omega = 2*pi*f", "变大：同样时间内振荡周期更多", "对车：助力节奏跟随更快步频"),
        ("alpha", "决定回到极限环的速度", "变大：扰动后更快恢复；过大可能数值敏感", "对车：节律更稳，但过大可能让助力变化突兀"),
        ("x", "主输出状态，用来调制阻尼", "x 正负变化让 D = D_base - lambda*x 周期变化", "对车：不是直接轮速，而是改变推行手感"),
        ("y", "辅助状态，与 x 构成二维相位", "提供旋转自由度，使系统形成稳定闭合轨道", "对车：保证节律连续，不只是一维摆动"),
    ]
    x0, y0, w, h = 920, 150, 790, 125
    colors = [(232, 239, 245), (245, 241, 232), (238, 245, 239), (242, 238, 248), (248, 238, 238)]
    for i, (name, line1, line2, line3) in enumerate(cards):
        y = y0 + i * (h + 18)
        d.rounded_rectangle((x0, y, x0 + w, y + h), radius=18, fill=colors[i], outline=(180, 190, 205), width=2)
        d.text((x0 + 26, y + 18), name, fill=(11, 37, 69), font=font_head)
        d.text((x0 + 120, y + 17), line1, fill=(36, 41, 47), font=font_body)
        d.text((x0 + 120, y + 52), line2, fill=(87, 96, 106), font=font_small)
        d.text((x0 + 120, y + 82), line3, fill=(87, 96, 106), font=font_small)

    d.text((70, 890), "工程实现：Hopf 产生节律状态 x/y；x 调制导纳阻尼 D；导纳模型把外力 F_ext 转换为车体速度 v_c；再换算为左右轮 rpm。", fill=(87, 96, 106), font=font_body)
    img.save(LIMIT_CYCLE_IMG)


def add_key_value_table(doc, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table)
    if widths is None:
        widths = [1.2, 2.1, 3.2]

    hdr = table.rows[0].cells
    for i, text in enumerate(rows[0]):
        hdr[i].text = ""
        set_cell_shading(hdr[i], "F2F4F7")
        set_cell_margins(hdr[i])
        p = hdr[i].paragraphs[0]
        r = p.add_run(text)
        set_run_font(r, size=10, bold=True, color="0B2545")
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        hdr[i].width = Inches(widths[i])

    for row in rows[1:]:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = ""
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            r = p.add_run(text)
            set_run_font(r, size=9.5)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[i].width = Inches(widths[i])
    doc.add_paragraph()
    return table


def build_doc():
    draw_limit_cycle()

    doc = Document()
    style_document(doc)

    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Hopf 振荡器参数详解")
    set_run_font(r, size=22, bold=True, color="0B2545")

    p = doc.add_paragraph(style="Subtitle")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("面向 STM32H7 智能助行车 CPG 变导纳控制的参数说明、极限环影响与使用体验分析")
    set_run_font(r, size=11, color="57606A")

    add_para(doc, "本文用于解释最初 Hopf 振荡器模型中的 x、y、mu、omega、alpha 五个核心参数。说明重点不是停留在数学符号，而是把参数含义、定义原因、对极限环的影响，以及它们最终如何改变助行车推行手感联系起来。")

    add_para(doc, "一、工程里的 Hopf 振荡器到底在做什么", style="Heading 1")
    add_para(doc, "在本助行车项目中，Hopf 振荡器不是直接生成轮速，也不是直接生成关节角度。它的作用是提供一个稳定、可调频、可调幅的周期节律状态。这个状态再进入导纳控制，用来动态调制阻尼，使助行车的助力特性能够跟随人的步态节奏变化。")
    add_formula(doc, "ẋ = α(μ − x² − y²)x − ωy + εF_ext", "x 方向状态方程：包含振幅收敛项、相位旋转项和外力耦合项。")
    add_formula(doc, "ẏ = α(μ − x² − y²)y + ωx", "y 方向状态方程：与 x 共同形成二维旋转系统。")
    add_formula(doc, "D_current = D_base − λx", "工程落地：用 x 调制导纳阻尼，而不是把 x 直接作为电机速度。")
    add_formula(doc, "M·v̇_c + D_current·v_c = F_ext", "导纳输出：外力通过虚拟质量和动态阻尼转换为车体线速度。")

    add_para(doc, "核心参数总览", style="Heading 2")
    add_key_value_table(doc, [
        ["参数", "数学意义", "在助行车中的工程意义"],
        ["x", "Hopf 主状态，参与二维相位旋转，是可输出的节律变量。", "用于调制阻尼 D，使助力强弱出现步态节律变化；不应直接等同轮速。"],
        ["y", "Hopf 辅助状态，与 x 组成二维振荡器。", "维持连续相位和稳定极限环，使系统不是简单一维来回摆动。"],
        ["mu", "目标振幅平方，稳定半径为 sqrt(mu)。", "控制节律调制幅度；越大，阻尼变化越明显。"],
        ["omega", "固有角频率，omega = 2*pi*f。", "控制节律快慢，可由动捕步频实时更新。"],
        ["alpha", "收敛速度，决定偏离极限环后恢复快慢。", "决定助力节律稳定性和扰动恢复速度。"],
    ], widths=[1.0, 2.8, 3.8])

    add_para(doc, "二、为什么 Hopf 振荡器要这样定义", style="Heading 1")
    add_para(doc, "Hopf 振荡器的关键设计是把“振幅控制”和“频率控制”分开。振幅由 mu 决定，频率由 omega 决定，收敛速度由 alpha 决定。这种结构非常适合助行车，因为人的步频会变化，但控制器不能因为步频变化就把助力幅度也一起扰乱。")
    add_formula(doc, "r² = x² + y²", "r 表示状态点到原点的距离，也就是振荡幅值。")
    add_formula(doc, "ṙ = α(μ − r²)r", "径向动力学：决定状态点向目标半径收敛。")
    add_formula(doc, "θ̇ = ω", "角向动力学：决定状态点沿极限环旋转的速度。")
    add_para(doc, "从这两个极坐标关系可以看出：mu 只影响极限环半径，omega 只影响沿环旋转的快慢，alpha 只影响回到环上的速度。这就是“振幅与频率解耦”的具体含义。")

    add_para(doc, "三、极限环是什么，以及这些参数如何影响极限环", style="Heading 1")
    add_para(doc, "极限环可以理解为系统最终会自动收敛到的一条闭合轨道。在 x-y 平面内，如果初始状态在环内，系统会向外扩张；如果初始状态在环外，系统会向内收缩；最后稳定在半径约为 sqrt(mu) 的圆形轨道上。")
    if LIMIT_CYCLE_IMG.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(LIMIT_CYCLE_IMG), width=Inches(6.5))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = cap.add_run("图 1  Hopf 参数对极限环和助行车体验的影响关系")
        set_run_font(rr, size=9.5, color="57606A")

    add_para(doc, "四、参数逐项详解", style="Heading 1")

    add_para(doc, "1. x：主振荡状态", style="Heading 2")
    add_para(doc, "x 是二维 Hopf 状态中的主状态。它会随着时间周期变化，并且在稳定后保持有限幅值。你最开始的资料里说 x 是“主要输出信号”，这个说法对腿式机器人或关节控制是成立的，但在助行车上需要改成更谨慎的理解：x 是主要节律信号，不是直接电机速度。")
    add_bullet(doc, "对极限环的影响：x 是极限环横轴坐标，随相位在 -sqrt(mu) 到 +sqrt(mu) 附近变化。")
    add_bullet(doc, "对算法的影响：当前工程用 x 计算 D_current = D_base - lambda*x。")
    add_bullet(doc, "对体验的影响：x 为正时阻尼降低，用户感觉更容易推动；x 为负时阻尼升高，用户感觉更稳、更不容易突然窜动。")
    add_bullet(doc, "注意：如果把 x 直接映射为轮速，轮子会前后往复，这不符合助行车实际使用逻辑。")

    add_para(doc, "2. y：辅助状态变量", style="Heading 2")
    add_para(doc, "y 不直接作为控制输出，但它非常关键。没有 y，系统只是一个一维非线性状态，很难自然形成连续相位和闭合极限环。y 的作用是和 x 共同构成二维相平面，使状态点可以围绕原点持续旋转。")
    add_bullet(doc, "对极限环的影响：y 是极限环纵轴坐标，和 x 一起决定 r^2 = x^2 + y^2。")
    add_bullet(doc, "对算法的影响：y 进入 dx、dy 的旋转项，保证 x 的变化不是突变，而是连续相位演化。")
    add_bullet(doc, "对体验的影响：助力变化更平滑，更像跟随人的周期步态，而不是简单的开关式助力。")

    add_para(doc, "3. mu：振幅参数", style="Heading 2")
    add_para(doc, "mu 是目标振幅的平方。稳定时 r = sqrt(mu)，所以 mu 不是直接等于振幅，而是振幅的平方。这样定义的好处是可以用同一个项 alpha(mu-r^2) 同时处理“振幅太小”和“振幅太大”两种情况。")
    add_formula(doc, "r_stable = √μ", "当前工程默认 mu = 1.0，因此稳定半径约为 1。")
    add_bullet(doc, "mu 增大：极限环半径变大，x 的正负摆幅变大。")
    add_bullet(doc, "mu 减小：极限环半径变小，x 的调制幅度变弱。")
    add_bullet(doc, "对助行车：mu 越大，D_current 的周期变化越明显，用户会感觉助力节律更强；mu 太大则可能让车的阻尼变化过强，推行体验不够稳。")
    add_bullet(doc, "工程建议：先保持 mu = 1.0，等实车数据稳定后再小范围调参。")

    add_para(doc, "4. omega：频率参数", style="Heading 2")
    add_para(doc, "omega 是角频率，决定状态点沿极限环旋转的快慢。它和普通频率 f 的关系是 omega = 2*pi*f。你的代码默认 f = 0.9 Hz，也就是助行器内部节律默认每秒约 0.9 个周期。")
    add_formula(doc, "ω = 2πf", "f 是步频，单位 Hz；omega 单位 rad/s。")
    add_bullet(doc, "omega 增大：极限环半径不变，但状态点绕环更快，x 的周期变化更快。")
    add_bullet(doc, "omega 减小：状态点绕环更慢，助力节奏变慢。")
    add_bullet(doc, "对助行车：omega 应该尽量跟随人的实际步频，否则助力节律会和人的步态错位。")
    add_bullet(doc, "工程实现：CPG_Update_Frequency() 接收 0.3 到 2.0 Hz 的步频输入，并更新 omega。")

    add_para(doc, "5. alpha：收敛速度参数", style="Heading 2")
    add_para(doc, "alpha 决定系统从当前状态回到极限环的速度。它不改变最终极限环半径，也不直接改变频率，而是改变“恢复速度”。在人推助行车时，外力输入、传感器噪声和起步停止都会扰动 Hopf 状态，因此 alpha 决定节律状态的稳定程度。")
    add_bullet(doc, "alpha 较大：系统更快回到极限环，节律更稳定，但数值积分更敏感。")
    add_bullet(doc, "alpha 较小：系统恢复更慢，节律更柔和，但受到扰动后可能拖泥带水。")
    add_bullet(doc, "对助行车：alpha 过小会让助力节律跟随慢；alpha 过大可能使阻尼变化显得硬，甚至在离散积分下出现抖动。")
    add_bullet(doc, "当前工程默认 alpha = 10.0，在 1 ms 控制周期下是一个便于先验证功能的保守值。")

    add_para(doc, "五、参数变化对极限环和车辆体验的对应关系", style="Heading 1")
    add_key_value_table(doc, [
        ["调参动作", "极限环变化", "车体体验变化"],
        ["增大 mu", "半径增大，x/y 振幅变大。", "阻尼调制更强，助力节律更明显；过大时可能突兀。"],
        ["减小 mu", "半径减小，振荡幅度变弱。", "助力变化更平缓，但 CPG 存在感变弱。"],
        ["增大 omega", "半径基本不变，绕环速度变快。", "适合快步频；过快时用户可能感觉助力节奏超前。"],
        ["减小 omega", "半径基本不变，绕环速度变慢。", "适合慢步频；过慢时可能跟不上人的步态。"],
        ["增大 alpha", "更快回到极限环。", "抗扰恢复强；过大可能让速度/阻尼变化显硬。"],
        ["减小 alpha", "回环速度变慢。", "更柔和；过小会让节律恢复慢、跟随不及时。"],
        ["直接使用 x 作轮速", "数学上能振荡，但输出有正负。", "轮子会前后交替，不符合助行车常规前进体验。"],
        ["使用 x 调制阻尼", "保留节律，同时不强迫速度正负振荡。", "更符合助行车：用户推力决定前进，CPG 改变助力手感。"],
    ], widths=[1.55, 2.45, 3.3])

    add_para(doc, "六、和当前代码的对应关系", style="Heading 1")
    add_para(doc, "当前工程在 CPG_Handle_t 中保存这些状态和参数，在 CPG_Step() 中每 1 ms 更新一次。对应关系如下。")
    add_key_value_table(doc, [
        ["工程变量", "默认值", "含义"],
        ["alpha", "10.0", "Hopf 收敛速度，影响扰动后回到极限环的快慢。"],
        ["mu", "1.0", "目标振幅平方，稳定半径约为 1。"],
        ["omega", "2*pi*0.9", "默认 0.9 Hz 步频节律，可由动捕步频更新。"],
        ["x", "0.1", "初始主状态，避免从原点完全静止启动。"],
        ["y", "0.0", "初始辅助状态。"],
        ["D_base", "40.0", "基础阻尼。"],
        ["lambda", "20.0", "x 到阻尼变化的映射增益。"],
        ["M", "15.0", "导纳模型虚拟质量。"],
        ["v_c", "0.0", "导纳输出车体线速度，后续换算为 rpm。"],
    ], widths=[1.35, 1.35, 4.6])
    add_para(doc, "代码路径：MDK-ARM/cpg_control.c 中完成 Hopf 状态更新、动态阻尼计算和导纳速度积分；MDK-ARM/Controller.c 中把 v_c 从 m/s 换算为轮速 rpm。")

    add_para(doc, "七、汇报时可以使用的讲解顺序", style="Heading 1")
    add_number(doc, "先说明问题：普通固定导纳控制只有固定 M、D，不能体现人的周期步态。")
    add_number(doc, "再说明引入 Hopf：Hopf 提供稳定极限环，能自然生成连续节律。")
    add_number(doc, "讲参数分工：mu 管振幅，omega 管频率，alpha 管恢复速度，x/y 管相位状态。")
    add_number(doc, "讲工程映射：不把 x 直接当轮速，而是用 x 调制阻尼 D。")
    add_number(doc, "讲体验结果：人的推力仍决定车前进，CPG 只让助力强弱随步态节律变化。")
    add_number(doc, "最后讲后续扩展：动捕步频通过 CPG_Update_Frequency() 更新 omega，实现人机节律同步。")

    if LOGIC_IMG.exists():
        doc.add_page_break()
        add_para(doc, "附录：当前工程 Hopf-CPG 变导纳控制流程图", style="Heading 1")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_picture(str(LOGIC_IMG), width=Inches(6.5))
        cap = doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rr = cap.add_run("图 2  Hopf-CPG 变导纳控制在工程中的数据流")
        set_run_font(rr, size=9.5, color="57606A")

    add_para(doc, "结论", style="Heading 1")
    add_para(doc, "Hopf 振荡器在本助行车项目中的价值，不是让车轮本身做正负振荡，而是提供一个稳定、可调频、可调幅的节律内核。mu 决定节律幅度，omega 决定节律快慢，alpha 决定节律稳定恢复能力，x/y 共同形成极限环。把 x 用于调制导纳阻尼，可以在不破坏助行车前进逻辑的前提下，让助力体验更接近人的自然步态节奏。")

    # Add a quiet footer.
    section = doc.sections[0]
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = footer.add_run("SmartRollator Hopf-CPG 参数说明")
    set_run_font(fr, size=8.5, color="57606A")

    doc.save(OUT)


if __name__ == "__main__":
    build_doc()
    print(OUT)
